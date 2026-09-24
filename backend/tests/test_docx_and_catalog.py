"""Тесты обогащения из DOCX и статистики каталога."""

from pathlib import Path

from app.models.robot import Robot
from app.services.import_docx_examples import (
    fuzzy_name_match,
    parse_payload_kg_from_fields,
    parse_ttx_fields,
)


FIXTURE_DIR = Path(__file__).parent / "fixtures"


def test_parse_ttx_fields_line_based():
    text = (
        "AMR (на примере модели Ronavi H1500)\n"
        "ТТХ:\n"
        "Грузоподъёмность: до 1500 кг.\n"
        "Максимальная скорость: до 1,5 м/с.\n"
        "Тип навигации: QR-метки.\n"
    )
    fields = parse_ttx_fields(text)
    assert "Грузоподъёмность" in fields
    assert "1500" in fields["Грузоподъёмность"]
    assert parse_payload_kg_from_fields(fields) == 1500.0


def test_fuzzy_name_contains():
    assert fuzzy_name_match("Ronavi H1500 (AMR)", "Ronavi H1500")
    assert fuzzy_name_match("Ronavi SR 50 кг", "Ronavi SR")
    assert fuzzy_name_match("DMR Carrier P", "DMR Carrier P")
    assert not fuzzy_name_match("DMR 300 Carrier B", "DMR Carrier P")
    assert not fuzzy_name_match("PuduBot 2", "Ronavi H1500")


def test_docx_import_with_mini_fixture(db_session, tmp_path):
    """Optional: if python-docx available, build tiny docx and enrich a robot."""
    try:
        from docx import Document
    except ImportError:
        return

    from app.services.import_docx_examples import import_docx_examples

    robot = Robot(
        name="Ronavi H1500 Test",
        payload_kg=None,
        price_rub=1_000_000,
        data_origin="source",
        confirmation_level="needs_review",
        archived=False,
        raw_data={},
    )
    db_session.add(robot)
    db_session.commit()
    db_session.refresh(robot)

    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(1, 0).text = (
        "AMR (на примере модели Ronavi H1500)\n"
        "ТТХ:\n"
        "Грузоподъёмность: до 1500 кг.\n"
        "Тип навигации: QR-метки.\n"
        "https://ronavi-robotics.ru/catalogue/h1500\n"
    )
    path = tmp_path / "Примеры_решений_типы_объектов.docx"
    doc.save(str(path))

    report = import_docx_examples(
        db_session, path=str(path), write_reports=False, commit=True
    )
    assert report.examples_found >= 1
    assert report.robots_updated >= 1
    db_session.refresh(robot)
    assert robot.payload_kg == 1500.0
    assert robot.raw_data and "docx_enrichment" in robot.raw_data
    assert robot.raw_data.get("payload_kg_source") == "docx"
    assert robot.source_url and "ronavi" in robot.source_url


def test_catalog_stats(client, db_session):
    db_session.add(
        Robot(
            name="Source Bot",
            data_origin="source",
            confirmation_level="confirmed",
            archived=False,
            price_rub=100,
        )
    )
    db_session.add(
        Robot(
            name="Demo Bot Extra",
            data_origin="demo",
            confirmation_level="assumption",
            archived=False,
        )
    )
    db_session.commit()

    r = client.get("/api/catalog/stats")
    assert r.status_code == 200
    data = r.json()
    assert "total" in data and "source_count" in data and "demo_count" in data
    assert data["total"] == data["source_count"] + data["demo_count"] or data["total"] >= data["source_count"]
    assert data["source_count"] >= 1
    assert data["demo_count"] >= 1


def test_catalog_filters_distinct_and_price(client, db_session):
    db_session.add(
        Robot(
            name="Filter Bot",
            data_origin="source",
            catalog_type_raw="brs",
            subtype_raw="AMR",
            industry_raw="склад",
            region_raw="РФ",
            scenario_raw="паллеты",
            availability_status="available",
            price_rub=500_000,
            confirmation_level="confirmed",
            archived=False,
        )
    )
    db_session.commit()

    meta = client.get("/api/catalog/meta/filters")
    assert meta.status_code == 200
    body = meta.json()
    assert "industry_raw" in body
    assert "region_raw" in body
    assert "subtype_raw" in body
    assert "scenario_raw" in body
    assert "catalog_type_raw" in body

    listed = client.get(
        "/api/catalog",
        params={"data_origin": "source", "min_price": 100_000, "max_price": 600_000, "page_size": 50},
    )
    assert listed.status_code == 200
    names = [i["name"] for i in listed.json()["items"]]
    assert "Filter Bot" in names
    item = next(i for i in listed.json()["items"] if i["name"] == "Filter Bot")
    assert "data_confidence" in item
