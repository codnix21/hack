"""Обогащение каталога из Примеры_решений_типы_объектов.docx (только чтение источника)."""

from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Optional

from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from app.models.robot import Robot

DOCX_FILENAME = "Примеры_решений_типы_объектов.docx"

MODEL_PATTERNS = [
    re.compile(
        r"(?:на примере(?: модели)?|модели|model)\s+([A-Za-zА-Яа-я0-9][A-Za-zА-Яа-я0-9\s\-/.]{1,40})",
        re.IGNORECASE,
    ),
    re.compile(
        r"\b(Ronavi\s+[A-Z0-9]+|DMR\s+Carrier\s+P|MARK\s+2\s+SE|PuduBot\s*2|"
        r"EVOCARGO\s+N1|Cognitive\s+Pilot|Pallet\s+[Ss]huttle|SmartCube|AMR100|Moros)\b"
    ),
]

PAYLOAD_KEY_RE = re.compile(r"грузоподъ[её]мност", re.IGNORECASE)
PAYLOAD_VAL_RE = re.compile(
    r"(?:до\s+)?([\d\s\u00a0]+(?:[.,]\d+)?)\s*(кг|т|тонн)?",
    re.IGNORECASE,
)
URL_RE = re.compile(r"https?://[^\s\)\"'<>]+")
KV_LINE_RE = re.compile(r"^([^:\n]{2,80})\s*:\s*(.+)$")


def _project_root() -> Path:
    # Local: backend/app/services/this.py → parents[3] = repo root
    # Docker: /app/app/services/this.py → parents[2] = /app (repo-like)
    here = Path(__file__).resolve()
    candidates = [
        here.parents[3],
        here.parents[2],
        Path("/data").parent if Path("/data").exists() else None,
        Path.cwd(),
    ]
    for root in candidates:
        if root is None:
            continue
        if (root / "data" / "sources").is_dir() or (root / "docs").is_dir():
            return root
        if (root / "data").is_dir():
            return root
    return here.parents[2]


def _docx_candidates(explicit: Optional[str] = None) -> list[Path]:
    if explicit:
        return [Path(explicit)]
    root = _project_root()
    dirs = [
        Path("/data/sources"),
        root / "data" / "sources",
        Path("data") / "sources",
        Path("..") / "data" / "sources",
    ]
    out: list[Path] = []
    for d in dirs:
        exact = d / DOCX_FILENAME
        out.append(exact)
        if d.is_dir():
            out.extend(sorted(d.glob("Примеры*.docx")))
            out.extend(sorted(d.glob("*решений*.docx")))
    # unique preserve order
    seen: set[str] = set()
    unique: list[Path] = []
    for p in out:
        key = str(p.resolve()) if p.exists() else str(p)
        if key not in seen:
            seen.add(key)
            unique.append(p)
    return unique


def resolve_docx_path(path: Optional[str] = None) -> Path:
    for p in _docx_candidates(path):
        if p.is_file():
            return p
    tried = ", ".join(str(p) for p in _docx_candidates(path)[:6])
    raise FileNotFoundError(f"DOCX не найден. Проверены: {tried}")


def _norm(s: str) -> str:
    s = (s or "").lower().replace("ё", "е")
    s = re.sub(r"[^\w\s]+", " ", s, flags=re.UNICODE)
    return re.sub(r"\s+", " ", s).strip()


def parse_ttx_fields(text: str) -> dict[str, str]:
    """Извлекает пары ключ→значение из блока ТТХ (только то, что есть в тексте)."""
    fields: dict[str, str] = {}
    if not text:
        return fields
    chunk = text.replace("\r\n", "\n").replace("\r", "\n")
    m = re.search(r"ТТХ\s*:\s*", chunk, re.IGNORECASE)
    body = chunk[m.end() :] if m else chunk

    for line in body.split("\n"):
        line = line.strip().strip(".;")
        if not line:
            continue
        km = KV_LINE_RE.match(line)
        if not km:
            continue
        key = km.group(1).strip(" .;-")
        val = km.group(2).strip(" .;-")
        if 2 <= len(key) <= 80 and val:
            fields[key] = val

    # Fallback: inline "Key: value Key2: value" after ТТХ
    if not fields and m:
        for km in re.finditer(
            r"([А-Яа-яA-Za-z][А-Яа-яA-Za-z0-9\-\s/]{1,50}?)\s*:\s*"
            r"([^:]+?)(?=\s+[А-Яа-яA-Za-z][^:]{0,40}:|$)",
            body.replace("\n", " "),
        ):
            key = km.group(1).strip(" .;-")
            val = km.group(2).strip(" .;-")
            if 2 <= len(key) <= 60 and val and len(val) < 200:
                fields[key] = val
    return fields


def extract_urls(text: str) -> list[str]:
    return URL_RE.findall(text or "")


def extract_name_hints(text: str) -> list[str]:
    hints: list[str] = []
    for pat in MODEL_PATTERNS:
        for m in pat.finditer(text or ""):
            hint = m.group(1).strip(" .;,)")
            hint = re.sub(r"\s+", " ", hint)
            if len(hint) >= 3 and hint not in hints:
                hints.append(hint)
    return hints


def parse_payload_kg_from_fields(fields: dict[str, str]) -> Optional[float]:
    for key, val in fields.items():
        if not PAYLOAD_KEY_RE.search(key):
            continue
        m = PAYLOAD_VAL_RE.search(val.replace(",", "."))
        if not m:
            continue
        num = m.group(1).replace(" ", "").replace("\u00a0", "")
        try:
            value = float(num)
        except ValueError:
            continue
        unit = (m.group(2) or "кг").lower()
        if unit.startswith("т"):
            value *= 1000.0
        return value
    return None


def fuzzy_name_match(robot_name: str, hint: str) -> bool:
    rn, hn = _norm(robot_name), _norm(hint)
    if not rn or not hn:
        return False
    if hn in rn or rn in hn:
        return True
    rn_base = rn.split("(")[0].strip()
    if hn in rn_base or rn_base in hn:
        return True
    # Все токены подсказки должны быть целыми словами имени (не частичный overlap)
    h_tokens = hn.split()
    r_tokens = set(rn_base.split())
    if len(h_tokens) >= 2 and all(t in r_tokens for t in h_tokens):
        return True
    return False


@dataclass
class DocxExample:
    name_hint: str
    fields: dict[str, str]
    urls: list[str]
    snippet: str
    table: Optional[int] = None
    row: Optional[int] = None
    col: Optional[int] = None


@dataclass
class DocxImportReport:
    source_file: str
    examples_found: int = 0
    robots_matched: int = 0
    robots_updated: int = 0
    payload_set: int = 0
    source_url_set: int = 0
    unmatched_hints: list[str] = field(default_factory=list)
    mappings: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def extract_examples_from_docx(path: Path) -> tuple[list[DocxExample], list[dict[str, Any]]]:
    from docx import Document  # lazy import

    doc = Document(str(path))
    examples: list[DocxExample] = []
    url_entries: list[dict[str, Any]] = []

    paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for i, text in enumerate(paras):
        urls = extract_urls(text)
        if urls:
            prev = paras[i - 1] if i else ""
            url_entries.append({"paragraph": i, "context": prev, "text": text, "urls": urls})

    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                text = (cell.text or "").strip()
                if not text or "ТТХ" not in text.upper():
                    continue
                fields = parse_ttx_fields(text)
                urls = extract_urls(text)
                hints = extract_name_hints(text)
                if not hints:
                    continue
                for hint in hints:
                    # attach nearby URLs by name
                    related = list(urls)
                    hn = _norm(hint)
                    for ue in url_entries:
                        ctx = _norm(ue.get("context", "") + " " + ue.get("text", ""))
                        if hn and hn in ctx:
                            for u in ue["urls"]:
                                if u not in related:
                                    related.append(u)
                    examples.append(
                        DocxExample(
                            name_hint=hint,
                            fields=fields,
                            urls=related,
                            snippet=text[:500],
                            table=ti,
                            row=ri,
                            col=ci,
                        )
                    )

    # Deduplicate by hint keeping richest fields
    by_hint: dict[str, DocxExample] = {}
    for ex in examples:
        key = _norm(ex.name_hint)
        prev = by_hint.get(key)
        if prev is None or len(ex.fields) > len(prev.fields) or (
            len(ex.fields) == len(prev.fields) and len(ex.urls) > len(prev.urls)
        ):
            by_hint[key] = ex
    return list(by_hint.values()), url_entries


def _write_mapping_reports(report: DocxImportReport) -> None:
    root = _project_root()
    imported_dir = root / "data" / "imported"
    imported_dir.mkdir(parents=True, exist_ok=True)

    payload = report.to_dict()
    (imported_dir / "docx_mapping.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    lines = [
        "# Соответствие примеров из DOCX каталогу",
        "",
        f"Источник: `{report.source_file}` (только чтение; оригиналы в `data/sources/` не изменяются).",
        "",
        f"- Найдено примеров: **{report.examples_found}**",
        f"- Сопоставлено с роботами: **{report.robots_matched}**",
        f"- Обновлено записей: **{report.robots_updated}**",
        f"- Установлено payload_kg из DOCX: **{report.payload_set}**",
        f"- Установлено source_url: **{report.source_url_set}**",
        "",
        "## Таблица сопоставления (DOCX ↔ CSV/БД)",
        "",
        "| Источник | Решение | Поле | Значение | Таблица БД | Статус сопоставления |",
        "|---|---|---|---|---|---|",
    ]
    for m in report.mappings:
        robot = m.get("robot_name") or m.get("name_hint") or ""
        loc = f"таблица {m.get('table')}, строка {m.get('row')}, ячейка {m.get('col')}"
        fields = m.get("fields") or {}
        if not fields:
            for key in m.get("field_keys") or []:
                lines.append(
                    f"| DOCX ({loc}) | {robot} | {key} | "
                    f"(см. raw_data.docx_enrichment) | robots.raw_data | "
                    f"{m.get('match_status', 'сопоставлено')} |"
                )
        else:
            for key, val in fields.items():
                safe_val = str(val).replace("|", "\\|").replace("\n", " ")
                db_col = "robots.payload_kg" if "грузоподъ" in key.lower() else "robots.raw_data.docx_enrichment"
                lines.append(
                    f"| DOCX ({loc}) | {robot} | {key} | {safe_val} | {db_col} | "
                    f"{m.get('match_status', 'сопоставлено')} |"
                )
        for u in m.get("urls") or []:
            lines.append(
                f"| DOCX ({loc}) | {robot} | source_url | {u} | robots.source_url | "
                f"{m.get('match_status', 'сопоставлено')} |"
            )
    if report.unmatched_hints:
        lines.extend(["", "## Не сопоставлено с CSV/каталогом", ""])
        for h in report.unmatched_hints:
            lines.append(f"- **{h}** — нет совпадения имени в каталоге CSV")
    if report.errors:
        lines.extend(["", "## Ошибки", ""])
        for e in report.errors:
            lines.append(f"- {e}")
    lines.extend(
        [
            "",
            "## Краткое сопоставление по моделям",
            "",
            "| Подсказка DOCX | Робот | ID | Поля ТТХ | URL |",
            "|---|---|---:|---|---|",
        ]
    )
    for m in report.mappings:
        field_list = ", ".join(m.get("field_keys") or []) or "—"
        urls = ", ".join(m.get("urls") or []) or "—"
        lines.append(
            f"| {m.get('name_hint','')} | {m.get('robot_name','')} | "
            f"{m.get('robot_id','')} | {field_list} | {urls} |"
        )
    lines.append("")
    (imported_dir / "docx_mapping.md").write_text("\n".join(lines), encoding="utf-8")


def import_docx_examples(
    db: Session,
    path: Optional[str] = None,
    *,
    write_reports: bool = True,
    commit: bool = True,
) -> DocxImportReport:
    """Обогащает matched robots: raw_data.docx_enrichment; опционально payload/source_url."""
    try:
        docx_path = resolve_docx_path(path)
    except FileNotFoundError as e:
        report = DocxImportReport(source_file=DOCX_FILENAME)
        report.errors.append(str(e))
        if write_reports:
            _write_mapping_reports(report)
        return report

    report = DocxImportReport(source_file=docx_path.name)
    try:
        examples, _url_entries = extract_examples_from_docx(docx_path)
    except Exception as e:  # noqa: BLE001
        report.errors.append(f"Ошибка чтения DOCX: {e}")
        if write_reports:
            _write_mapping_reports(report)
        return report

    report.examples_found = len(examples)
    robots = db.query(Robot).filter(Robot.archived.is_(False)).all()

    matched_robot_ids: set[int] = set()
    for ex in examples:
        matches = [r for r in robots if fuzzy_name_match(r.name, ex.name_hint)]
        if not matches:
            report.unmatched_hints.append(ex.name_hint)
            continue

        for robot in matches:
            matched_robot_ids.add(robot.id)
            raw = dict(robot.raw_data or {})
            enrichment = {
                "fields": dict(ex.fields),
                "source_file": docx_path.name,
                "urls": list(ex.urls),
                "matched_snippet": ex.snippet,
                "name_hint": ex.name_hint,
            }
            raw["docx_enrichment"] = enrichment
            robot.raw_data = raw
            flag_modified(robot, "raw_data")

            if robot.payload_kg is None:
                payload = parse_payload_kg_from_fields(ex.fields)
                if payload is not None:
                    robot.payload_kg = payload
                    notes = dict(raw.get("field_sources") or {})
                    notes["payload_kg"] = "docx"
                    raw["field_sources"] = notes
                    raw["payload_kg_source"] = "docx"
                    robot.raw_data = raw
                    flag_modified(robot, "raw_data")
                    report.payload_set += 1

            if not robot.source_url and ex.urls:
                robot.source_url = ex.urls[0]
                report.source_url_set += 1

            if robot.data_origin == "demo":
                pass  # keep origin; enrichment is additive
            elif robot.data_origin == "source":
                # mark lightly enriched without inventing origin rewrite unless already source
                pass

            report.robots_updated += 1
            report.mappings.append(
                {
                    "name_hint": ex.name_hint,
                    "robot_id": robot.id,
                    "robot_name": robot.name,
                    "field_keys": list(ex.fields.keys()),
                    "fields": dict(ex.fields),
                    "urls": ex.urls,
                    "table": ex.table,
                    "row": ex.row,
                    "col": ex.col,
                    "match_status": "сопоставлено",
                }
            )

    report.robots_matched = len(matched_robot_ids)

    if commit:
        db.commit()

    if write_reports:
        _write_mapping_reports(report)
    return report
