"""Шаблоны импорта параметров объекта (Excel + маппинг русских названий)."""

from __future__ import annotations

import io
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ключ → (название RU, пример, единица, подсказка)
WAREHOUSE_TEMPLATE: list[tuple[str, str, str, str, str]] = [
    ("area_width_m", "Ширина площадки", "40", "м", "Габарит зоны по ширине"),
    ("area_length_m", "Длина площадки", "60", "м", "Габарит зоны по длине"),
    ("aisle_width_mm", "Ширина проходов", "2500", "мм", "Минимальная ширина прохода для робота"),
    ("required_payload_kg", "Требуемая грузоподъёмность", "500", "кг", "Максимальная масса груза"),
    ("peak_demand", "Пиковый спрос", "120", "оп/ч", "Пиковая интенсивность операций в час"),
    ("zones_count", "Число рабочих зон", "4", "шт.", "Количество зон на объекте"),
    ("charging_stations", "Зарядных станций", "2", "шт.", "Точки зарядки AMR"),
    ("current_labor_cost_year", "Стоимость персонала", "4800000", "₽/год", "Годовые затраты на персонал без роботизации"),
    ("temperature_min_c", "Температура мин.", "5", "°C", "Нижняя граница рабочей температуры"),
    ("temperature_max_c", "Температура макс.", "25", "°C", "Верхняя граница рабочей температуры"),
    ("storage_type", "Тип хранения", "Паллетное", "", "Варианты: Паллетное / Полочное / Смешанное"),
]

AIRPORT_TEMPLATE: list[tuple[str, str, str, str, str]] = [
    ("area_width_m", "Ширина зоны", "80", "м", "Габарит зоны операции"),
    ("area_length_m", "Длина зоны", "120", "м", "Габарит зоны операции"),
    ("aisle_width_mm", "Ширина проездов", "3500", "мм", "Ширина проезда техники"),
    ("required_payload_kg", "Масса перемещаемых объектов", "300", "кг", "Типовая масса единицы"),
    ("peak_demand", "Пиковая нагрузка", "90", "оп/ч", "Пик операций в час"),
    ("zones_count", "Число зон операции", "3", "шт.", ""),
    ("charging_stations", "Зарядных станций", "2", "шт.", ""),
    ("current_labor_cost_year", "Стоимость персонала", "6200000", "₽/год", ""),
    ("zone_type", "Тип зоны", "Багажный комплекс", "", "Варианты: Перрон / Терминал / Багажный комплекс / Смешанная"),
]

MEDICAL_TEMPLATE: list[tuple[str, str, str, str, str]] = [
    ("area_width_m", "Ширина корпуса", "35", "м", ""),
    ("area_length_m", "Длина корпуса", "55", "м", ""),
    ("aisle_width_mm", "Ширина коридоров", "1800", "мм", ""),
    ("required_payload_kg", "Масса грузовой единицы", "80", "кг", ""),
    ("peak_demand", "Количество перевозок", "40", "оп/ч", "Пик перевозок в час"),
    ("zones_count", "Число отделений / зон", "6", "шт.", ""),
    ("floors_count", "Этажность", "3", "эт.", ""),
    ("charging_stations", "Зарядных станций", "2", "шт.", ""),
    ("current_labor_cost_year", "Стоимость персонала", "3500000", "₽/год", ""),
    ("humidity_max_pct", "Влажность макс.", "60", "%", ""),
    ("facility_type", "Тип учреждения", "Больница", "", "Варианты: Больница / Клиника / Лаборатория"),
]

BY_CODE: dict[str, list[tuple[str, str, str, str, str]]] = {
    "warehouse": WAREHOUSE_TEMPLATE,
    "склад": WAREHOUSE_TEMPLATE,
    "airport": AIRPORT_TEMPLATE,
    "аэропорт": AIRPORT_TEMPLATE,
    "medical": MEDICAL_TEMPLATE,
    "медицина": MEDICAL_TEMPLATE,
    "медицинское": MEDICAL_TEMPLATE,
}

# Русские / альтернативные названия → канонический ключ
LABEL_TO_KEY: dict[str, str] = {}
ENUM_VALUE_MAP: dict[str, dict[str, str]] = {
    "storage_type": {
        "паллетное": "pallet",
        "pallet": "pallet",
        "полочное": "shelf",
        "shelf": "shelf",
        "смешанное": "mixed",
        "mixed": "mixed",
    },
    "zone_type": {
        "перрон": "apron",
        "apron": "apron",
        "терминал": "terminal",
        "terminal": "terminal",
        "багажный комплекс": "baggage",
        "baggage": "baggage",
        "смешанная": "mixed",
        "mixed": "mixed",
    },
    "facility_type": {
        "больница": "hospital",
        "hospital": "hospital",
        "клиника": "clinic",
        "clinic": "clinic",
        "лаборатория": "lab",
        "lab": "lab",
    },
}


def _register_labels() -> None:
    for rows in BY_CODE.values():
        for key, label, *_rest in rows:
            LABEL_TO_KEY[label.strip().lower()] = key
            LABEL_TO_KEY[key.strip().lower()] = key
    # доп. алиасы
    extras = {
        "площадь": "area_m2",
        "ширина": "area_width_m",
        "длина": "area_length_m",
        "грузоподъёмность": "required_payload_kg",
        "грузоподъемность": "required_payload_kg",
        "пиковый спрос": "peak_demand",
        "число зон": "zones_count",
        "зарядных станций": "charging_stations",
        "стоимость персонала": "current_labor_cost_year",
    }
    LABEL_TO_KEY.update(extras)


_register_labels()


def resolve_param_key(raw: str) -> str:
    """Русское название или ключ → канонический ключ."""
    s = str(raw or "").strip()
    if not s:
        return s
    low = s.lower()
    if low in LABEL_TO_KEY:
        return LABEL_TO_KEY[low]
    # «Ширина площадки, м» → без единицы
    if "," in low:
        base = low.split(",", 1)[0].strip()
        if base in LABEL_TO_KEY:
            return LABEL_TO_KEY[base]
    return s


def normalize_param_value(key: str, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        enum_map = ENUM_VALUE_MAP.get(key)
        if enum_map and text.lower() in enum_map:
            return enum_map[text.lower()]
        return text
    return value


def resolve_template_rows(
    object_type_code: Optional[str] = None,
    parameter_schema: Optional[dict] = None,
) -> list[tuple[str, str, str, str, str]]:
    code = (object_type_code or "warehouse").strip().lower()
    if code in BY_CODE:
        return BY_CODE[code]
    props = (parameter_schema or {}).get("properties") or {}
    if props:
        rows: list[tuple[str, str, str, str, str]] = []
        for key, meta in props.items():
            if not isinstance(meta, dict):
                continue
            title = str(meta.get("title") or key)
            unit = str(meta.get("unit") or "")
            sample = ""
            if meta.get("type") == "number":
                sample = "0"
            elif meta.get("enum"):
                sample = str(meta["enum"][0])
            rows.append((key, title, sample, unit, ""))
        if rows:
            return rows
    return WAREHOUSE_TEMPLATE


def object_type_title(object_type_code: Optional[str] = None) -> str:
    code = (object_type_code or "warehouse").strip().lower()
    if code in ("airport", "аэропорт"):
        return "Аэропорт"
    if code in ("medical", "медицина", "медицинское"):
        return "Медицинское учреждение"
    return "Склад"


def build_params_template_xlsx(
    object_type_code: Optional[str] = None,
    parameter_schema: Optional[dict] = None,
) -> bytes:
    """Excel-шаблон: лист Инструкция + Параметры (чистая таблица для импорта)."""
    rows = resolve_template_rows(object_type_code, parameter_schema)
    title = object_type_title(object_type_code)

    wb = Workbook()

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="14171C")
    hint_fill = PatternFill("solid", fgColor="F4F5F7")
    thin = Border(
        left=Side(style="thin", color="D1D5DE"),
        right=Side(style="thin", color="D1D5DE"),
        top=Side(style="thin", color="D1D5DE"),
        bottom=Side(style="thin", color="D1D5DE"),
    )
    value_fill = PatternFill("solid", fgColor="FFF7ED")

    # Instruction first
    wi = wb.active
    wi.title = "Инструкция"
    wi["A1"] = f"Шаблон параметров объекта — {title}"
    wi["A1"].font = Font(name="Calibri", bold=True, size=14, color="14171C")
    instructions = [
        "",
        "Как заполнить:",
        "1. Откройте лист «Параметры».",
        "2. В жёлтой колонке «Значение» замените примеры на данные вашего объекта.",
        "3. Названия в колонке «Параметр» не меняйте — по ним система распознаёт поля.",
        "4. Числа пишите без пробелов: 4800000 (не 4 800 000).",
        "5. Тип хранения / зоны / учреждения можно указать по-русски (см. колонку «Подсказка»).",
        "6. Сохраните файл и загрузите на шаге «Импорт» в проекте.",
        "",
        "Не удаляйте строки параметров, даже если значение пока неизвестно — оставьте ячейку пустой.",
    ]
    for i, line in enumerate(instructions, start=2):
        cell = wi.cell(row=i, column=1, value=line)
        cell.font = Font(name="Calibri", size=11, color="333843")
    wi.column_dimensions["A"].width = 95

    # Clean data sheet — header row 1 for pandas
    ws = wb.create_sheet("Параметры")
    headers = ["Параметр", "Значение", "Ед. изм.", "Подсказка"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, (_key, label, sample, unit, hint) in enumerate(rows):
        r = 2 + i
        c1 = ws.cell(row=r, column=1, value=label)
        c2 = ws.cell(row=r, column=2, value=_excel_value(sample))
        c3 = ws.cell(row=r, column=3, value=unit or "—")
        c4 = ws.cell(row=r, column=4, value=hint or "")
        for c in (c1, c2, c3, c4):
            c.border = thin
            c.font = Font(name="Calibri", size=11)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        c1.fill = hint_fill
        c2.fill = value_fill
        c3.fill = hint_fill
        c4.fill = hint_fill

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 52
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{1 + len(rows)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _excel_value(sample: str):
    if sample == "":
        return ""
    try:
        if "." in sample:
            return float(sample)
        return int(sample)
    except ValueError:
        return sample


def template_filename(object_type_code: Optional[str] = None) -> str:
    code = (object_type_code or "warehouse").strip().lower()
    if code in ("airport", "аэропорт"):
        return "shablon_parametrov_aeroport.xlsx"
    if code in ("medical", "медицина", "медицинское"):
        return "shablon_parametrov_medicina.xlsx"
    return "shablon_parametrov_sklad.xlsx"


# backward-compatible CSV (редко нужен)
def build_params_template_csv(
    object_type_code: Optional[str] = None,
    parameter_schema: Optional[dict] = None,
) -> str:
    rows = resolve_template_rows(object_type_code, parameter_schema)
    lines = ["Параметр,Значение,Ед. изм.,Подсказка"]
    for _key, label, sample, unit, hint in rows:
        safe = hint.replace('"', "'")
        lines.append(f'"{label}",{sample},"{unit}","{safe}"')
    return "\n".join(lines) + "\n"
