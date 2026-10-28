"""Экспорт проекта в PDF и Excel."""

from __future__ import annotations

import io
import os
from datetime import datetime
from typing import Any, Callable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# Visual identity for PDF (print-friendly, not dark UI)
_PDF_ORANGE = colors.Color(0.94, 0.45, 0.09)  # #f07316
_PDF_INK = colors.Color(0.08, 0.09, 0.12)
_PDF_MUTED = colors.Color(0.35, 0.38, 0.45)
_PDF_HEADER = colors.Color(0.12, 0.15, 0.20)
_PDF_ROW_ALT = colors.Color(0.97, 0.97, 0.98)
_PDF_CARD = colors.Color(0.98, 0.96, 0.93)
_PDF_CARD_LINE = colors.Color(0.94, 0.45, 0.09)


_FONT_REGISTERED = False
_FONT_NAME = "Helvetica"

PLATFORM_NAME = "Платформа подбора роботизированных решений"

STATUS_LABELS_RU = {
    "suitable": "Предварительно подходит",
    "needs_review": "Требует проверки",
    "excluded": "Не подходит",
}

METRIC_LABELS_RU = {
    "capex": "Капитальные затраты",
    "opex": "Эксплуатационные расходы",
    "opex_annual": "Эксплуатационные расходы",
    "roi": "Рентабельность инвестиций",
    "roi_percent": "Рентабельность инвестиций",
    "tco": "Совокупная стоимость владения",
    "npv": "Чистая приведённая стоимость",
    "irr": "Внутренняя норма доходности",
    "payback": "Срок окупаемости",
    "annual_effect": "Годовой экономический эффект",
    "cf": "Денежный поток",
    "straight_line": "Линейный метод",
}

WORK_MODE_LABELS_RU = {
    "1shift": "1 смена",
    "2shift": "2 смены",
    "24x7": "Круглосуточно (24/7)",
    "24/7": "Круглосуточно (24/7)",
}

ASSUMPTION_LABELS_RU = {
    "utilization": "Загрузка оборудования",
    "availability": "Доступность",
    "reserve": "Резерв парка",
    "years": "Горизонт анализа",
    "infra_share_of_equipment": "Доля инфраструктуры от оборудования",
    "integration_share": "Доля интеграции",
    "commissioning_share": "Доля пусконаладки",
    "training_per_robot": "Обучение на робота",
    "contingency_share": "Доля непредвиденных расходов",
    "energy_per_robot_year": "Энергия на робота",
    "connectivity_per_robot_year": "Связь на робота",
    "consumables_per_robot_year": "Расходники на робота",
    "repair_share_of_price": "Доля ремонта от цены",
    "ops_staff_per_10_robots": "Персонал на 10 роботов",
    "ops_staff_salary_year": "Зарплата персонала",
    "license_share_of_software": "Доля лицензий от ПО",
    "current_labor_cost_year": "Годовые затраты на персонал",
    "labor_reduction_share": "Ожидаемая доля сокращения трудозатрат",
    "additional_revenue": "Дополнительная выручка",
    "raas_monthly_per_robot": "Аренда в месяц на робота",
    "unit_price": "Цена решения",
    "peak_demand": "Пиковая нагрузка",
    "productivity": "Производительность",
    "robots_count": "Количество роботов",
    "software_cost": "Стоимость ПО",
    "maintenance_year": "Обслуживание",
    "implementation_cost": "Внедрение",
    "discount_rate": "Ставка дисконтирования",
    "battery_cycle_years": "Цикл замены АКБ",
    "battery_cost_share": "Доля стоимости АКБ от цены",
    "charging_station_cost": "Стоимость зарядной станции",
    "charging_stations_count": "Число зарядных станций",
    "depreciation_enabled": "Учёт амортизации",
    "depreciation_method": "Метод амортизации",
    "depreciation_years": "Срок амортизации",
    "salvage_value": "Ликвидационная стоимость",
    "assumed_unit_price": "Принятая цена решения",
}

# label, unit (optional)
OBJECT_PARAM_META: dict[str, tuple[str, str]] = {
    "area_m2": ("Площадь", "м²"),
    "active_area_m2": ("Площадь активной зоны", "м²"),
    "apron_area_m2": ("Площадь перрона", "м²"),
    "ceiling_height_m": ("Высота потолков", "м"),
    "floors_count": ("Число этажей", "шт"),
    "zones_count": ("Число зон", "шт"),
    "work_mode": ("Режим работы", ""),
    "shifts_count": ("Число смен", "шт"),
    "shift_hours": ("Продолжительность смены", "ч"),
    "staff_count": ("Численность персонала", "чел"),
    "staff_cost_year": ("Затраты на персонал", "₽/год"),
    "staff_cost_per_person_year": ("Стоимость сотрудника", "₽/год"),
    "ops_count": ("Операций в сутки", "оп/сут"),
    "peak_demand": ("Пиковый спрос", "оп/ч"),
    "peak_ops_per_hour": ("Пиковый спрос", "оп/ч"),
    "avg_demand": ("Средний спрос", "оп/ч"),
    "peak_factor": ("Пиковый коэффициент", ""),
    "sku_count": ("Число SKU", "шт"),
    "required_payload_kg": ("Требуемая грузоподъёмность", "кг"),
    "cargo_length_mm": ("Длина груза", "мм"),
    "cargo_width_mm": ("Ширина груза", "мм"),
    "cargo_height_mm": ("Высота груза", "мм"),
    "aisle_width_mm": ("Ширина прохода", "мм"),
    "aisle_width_m": ("Ширина прохода", "м"),
    "main_aisle_width_m": ("Ширина главных проездов", "м"),
    "route_length_m": ("Длина маршрута", "м"),
    "floor_type": ("Тип покрытия пола", ""),
    "temperature_min_c": ("Мин. температура", "°C"),
    "temperature_max_c": ("Макс. температура", "°C"),
    "humidity_max_pct": ("Макс. влажность", "%"),
    "charging_stations": ("Зарядные станции", "шт"),
    "load_points": ("Точки погрузки", "шт"),
    "unload_points": ("Точки разгрузки", "шт"),
    "layout_constraints": ("Ограничения планировки", ""),
    "area_width_m": ("Ширина площадки", "м"),
    "area_length_m": ("Длина площадки", "м"),
    "current_labor_cost_year": ("Текущие затраты на труд", "₽/год"),
    "required_processes": ("Требуемые процессы", ""),
    "indoor": ("Работа в помещении", ""),
    "required_navigation": ("Требуемая навигация", ""),
    "passenger_flow": ("Пассажиропоток", "чел/сут"),
    "cargo_flow": ("Грузопоток", "т/сут"),
    "access_constraints": ("Ограничения доступа", ""),
    "security_requirements": ("Требования безопасности", ""),
    "zone_type": ("Тип зоны", ""),
    "transport_count": ("Транспортировок в сутки", "шт"),
    "meds_share": ("Доля медикаментов", "%"),
    "meals_share": ("Доля питания", "%"),
    "linen_share": ("Доля белья", "%"),
    "waste_share": ("Доля отходов", "%"),
    "elevators": ("Лифты", "шт"),
    "sanitary_constraints": ("Санитарные ограничения", ""),
    "facility_type": ("Тип учреждения", ""),
    "terminals_count": ("Количество терминалов", "шт"),
    "gates_count": ("Количество гейтов", "шт"),
    "productivity": ("Производительность", "оп/ч"),
    "utilization": ("Загрузка", ""),
    "availability": ("Доступность", ""),
    "reserve": ("Резерв", ""),
    "years": ("Горизонт анализа", "лет"),
    "staff_salary_month": ("Средняя зарплата", "₽/мес"),
    "payroll_tax_coef": ("Коэффициент начислений на ФОТ", ""),
    "shifts_per_day": ("Число смен в сутки", "шт"),
    "hours_per_shift": ("Часов в смене", "ч"),
    # Датасеты_хакатон (транслит-ключи)
    "rovnost_pola_otklonenie": ("Ровность пола (отклонение)", "мм"),
    "rabochih_dney_v_godu": ("Рабочих дней в году", "дн"),
    "obem_priemki_poddony_sutki": ("Объём приёмки", "поддоны/сут"),
    "obem_otgruzki_poddony_sutki": ("Объём отгрузки", "поддоны/сут"),
    "obem_otbora_strok_sutki_vsego": ("Объём отбора", "строк/сут"),
    "obem_otbora_shtuk_sutki_vsego": ("Объём отбора", "шт/сут"),
    "dolya_melkoshtuchnogo_otbora_piece_pick": ("Доля мелкоштучного отбора", ""),
    "kolichestvo_sku_aktivnyh": ("Количество активных SKU", "шт"),
    "dolya_sku_s_bystrym_oborotom_a_klass": ("Доля SKU A-класса", ""),
    "obschaya_chislennost_personala_sklada": ("Общая численность персонала склада", "чел"),
    "iz_nih_otborschiki_komplektovschiki": ("Из них: отборщики", "чел"),
    "iz_nih_operatory_pogruzchikov": ("Из них: операторы погрузчиков", "чел"),
    "iz_nih_operatory_upakovochnyh_liniy": ("Из них: операторы упаковки", "чел"),
    "srednyaya_z_p_otborschika_gross": ("Средняя з/п отборщика", "₽"),
    "srednyaya_z_p_operatora_pogruzchika_gross": ("Средняя з/п оператора погрузчика", "₽"),
    "koeffitsient_nachisleniy_na_fot_strahovye_vznosy": ("Коэффициент начислений на ФОТ", ""),
    "srednyaya_vyrabotka_otborschika_strok_ch": ("Средняя выработка отборщика", "строк/ч"),
    "koeffitsient_poter_rabochego_vremeni_otpusk_bolezn_tekuchest": (
        "Коэффициент потерь рабочего времени",
        "",
    ),
    "srednyaya_dlina_marshruta_otborschika_na_1_stroku": ("Средняя длина маршрута отборщика", "м"),
    "protyazhennost_konveyernoy_transportnoy_sistemy": ("Протяжённость конвейерной системы", "м"),
    "tip_stellazhnoy_sistemy": ("Тип стеллажной системы", ""),
    "kolichestvo_palletomest": ("Количество паллетомест", "шт"),
    "srednyaya_massa_gruzovoy_edinitsy_pallet": ("Средняя масса паллеты", "кг"),
    "srednyaya_massa_shtuchnoy_edinitsy_sku": ("Средняя масса штучной единицы", "кг"),
}

ASSUMPTION_PARAM_KEYS = set(ASSUMPTION_LABELS_RU.keys()) | {
    "utilization",
    "availability",
    "reserve",
    "years",
}

PROJECT_FIELD_LABELS = {
    "id": "№",
    "name": "Название",
    "region": "Регион",
    "work_mode": "Режим работы",
    "status": "Статус",
    "description": "Описание",
    "is_demo": "Пример проекта",
    "object_type": "Тип объекта",
    "industry": "Отрасль",
    "created_at": "Создан",
    "updated_at": "Обновлён",
}

ECONOMICS_DESCRIPTIONS = {
    "robots_count": "Расчётное количество роботов для пикового спроса",
    "capex": "Первоначальные капитальные затраты",
    "opex_annual": "Эксплуатационные затраты в год",
    "annual_effect": "Ожидаемый чистый эффект за год",
    "payback_years": "Ориентировочный период возврата инвестиций",
    "roi_percent": "Возврат инвестиций относительно капитальных затрат",
    "tco": "Совокупная стоимость владения (капитальные + эксплуатационные × горизонт; без оставшегося ручного труда)",
}

BREAKDOWN_LABELS_RU = {
    "equipment": "Оборудование",
    "infrastructure": "Инфраструктура",
    "software": "Программное обеспечение",
    "integration": "Интеграция",
    "commissioning": "Пусконаладка",
    "training": "Обучение",
    "contingency": "Резерв / непредвиденные",
    "service": "Сервис и обслуживание",
    "licenses": "Лицензии",
    "energy": "Энергия",
    "connectivity": "Связь",
    "consumables": "Расходники",
    "repair": "Ремонт",
    "ops_staff": "Эксплуатационный персонал",
    "ops_staff_count": "Численность эксплуатационного персонала",
}

FORMULA_LABELS_RU = {
    "robots_count": "Количество роботов",
    "количество_роботов": "Количество роботов",
    "capex": "Капитальные затраты",
    "opex": "Эксплуатационные затраты",
    "annual_effect": "Годовой эффект",
    "payback": "Срок окупаемости",
    "roi": "Рентабельность инвестиций",
    "roi_horizon": "Рентабельность за горизонт анализа",
    "tco": "Совокупная стоимость владения (без оставшегося ручного труда)",
    "npv": "Чистая приведённая стоимость",
    "irr": "Внутренняя норма доходности",
    "gross_effect": "Валовой операционный эффект",
    "net_effect": "Чистый годовой эффект",
    "depreciation": "Амортизация капитальных затрат",
    "annual_effect_with_depreciation": "Годовой эффект с учётом амортизации",
}

PROJECT_STATUS_RU = {
    "draft": "Черновик",
    "active": "Активный",
    "archived": "Архив",
    "completed": "Завершён",
    "demo": "Каталог платформы",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=16, color="1F4E79")
SUBTITLE_FONT = Font(bold=True, size=12, color="2F5496")
SECTION_FILL = PatternFill("solid", fgColor="D6E3F0")
ALT_ROW_FILL = PatternFill("solid", fgColor="F5F8FB")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)

COMPARISON_FIELDS = [
    ("status", "Статус соответствия"),
    ("match_score", "Оценка, %"),
    ("data_origin", "Источник данных"),
]


def _register_cyrillic_font() -> str:
    global _FONT_REGISTERED, _FONT_NAME
    if _FONT_REGISTERED:
        return _FONT_NAME
    candidates = [
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
    ]
    for path in candidates:
        if os.path.isfile(path):
            try:
                pdfmetrics.registerFont(TTFont("AppCyrillic", path))
                _FONT_NAME = "AppCyrillic"
                _FONT_REGISTERED = True
                return _FONT_NAME
            except Exception:
                continue
    _FONT_REGISTERED = True
    return _FONT_NAME


def _scenario_metrics(results: Any) -> dict[str, Any]:
    if not isinstance(results, dict):
        return {}
    capex = results.get("capex") if isinstance(results.get("capex"), dict) else {}
    opex = results.get("opex") if isinstance(results.get("opex"), dict) else {}
    effect = results.get("annual_effect") if isinstance(results.get("annual_effect"), dict) else {}
    payback = results.get("payback") if isinstance(results.get("payback"), dict) else {}
    roi = results.get("roi") if isinstance(results.get("roi"), dict) else {}
    tco = results.get("tco") if isinstance(results.get("tco"), dict) else {}
    npv = results.get("npv") if isinstance(results.get("npv"), dict) else {}
    irr = results.get("irr") if isinstance(results.get("irr"), dict) else {}
    return {
        "robots_count": results.get("robots_count"),
        "capex": capex.get("capex") if capex else results.get("capex"),
        "opex_annual": opex.get("opex_annual") if opex else results.get("opex_annual"),
        "annual_effect": effect.get("annual_effect") if effect else results.get("annual_effect"),
        "gross_operating_effect": effect.get("gross_operating_effect"),
        "net_annual_effect": effect.get("net_annual_effect") or effect.get("annual_effect"),
        "payback_years": payback.get("payback_years") if payback else results.get("payback_years"),
        "roi_percent": roi.get("roi_percent") if roi else results.get("roi_percent"),
        "tco": tco.get("tco") if tco else results.get("tco"),
        "npv": npv.get("npv") if npv else results.get("npv"),
        "irr": irr.get("irr") if irr else results.get("irr"),
        "cashflows": results.get("cashflows") if isinstance(results.get("cashflows"), list) else [],
        "scenario_name_ru": results.get("scenario_name_ru") or results.get("scenario"),
    }


def _origin_label(origin: Optional[str], is_demo: bool = False) -> str:
    if origin == "source":
        return "Исходные материалы"
    if origin == "enriched":
        return "Обогащённые данные"
    if is_demo or origin == "demo" or not origin:
        return "Каталог платформы"
    return str(origin)


def _na(value: Any, empty: str = "Нет данных") -> Any:
    if value is None:
        return empty
    if isinstance(value, str) and not value.strip():
        return empty
    if isinstance(value, (list, dict)) and not value:
        return empty
    return value


def _dash(value: Any) -> Any:
    return _na(value, "—")


_VALUE_LABELS_RU = {
    "straight_line": "Линейный метод",
    "declining_balance": "Уменьшаемого остатка",
    "warehouse": "Склад",
    "airport": "Аэропорт",
    "medical": "Медицинское учреждение",
    "purchase": "Покупка",
    "raas": "Аренда",
    "both": "Покупка и аренда",
}


def _fmt_cell(value: Any) -> Any:
    if value is None:
        return "Нет данных"
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, float):
        if abs(value) >= 1000:
            return round(value, 2)
        if 0 < abs(value) < 1:
            return round(value, 4)
        return round(value, 2)
    if isinstance(value, (list, dict)):
        if not value:
            return "Нет данных"
        if isinstance(value, list):
            parts = []
            for item in value[:8]:
                if isinstance(item, dict):
                    parts.append(str(item.get("title") or item.get("name") or item.get("result") or item))
                else:
                    parts.append(str(item))
            return "; ".join(parts)
        return "Нет данных в исходных материалах"
    if isinstance(value, str) and value in _VALUE_LABELS_RU:
        return _VALUE_LABELS_RU[value]
    if isinstance(value, str) and value in METRIC_LABELS_RU:
        return METRIC_LABELS_RU[value]
    return value


def _format_money(value: Any) -> Any:
    if value is None or value == "":
        return "—"
    try:
        return round(float(value), 0)
    except (TypeError, ValueError):
        return value


def _money_label(value: Any) -> str:
    raw = _format_money(value)
    if raw == "—":
        return "—"
    try:
        return f"{int(round(float(raw))):,}".replace(",", " ") 
    except (TypeError, ValueError):
        return str(raw)


def _format_share_or_number(key: str, value: Any) -> Any:
    if value is None or value == "":
        return "—"
    try:
        num = float(value)
    except (TypeError, ValueError):
        return value
    if key.endswith("_share") or key in {
        "utilization",
        "availability",
        "reserve",
        "labor_reduction_share",
        "meds_share",
        "meals_share",
        "linen_share",
        "waste_share",
    }:
        # model stores 0..1 shares
        if 0 <= num <= 1:
            return round(num * 100, 1)
    if key.endswith("_year") or key.endswith("_rub") or "cost" in key or "price" in key or key.startswith("training"):
        return round(num, 0)
    return round(num, 2) if isinstance(value, float) else value


def _share_unit(key: str) -> str:
    if key.endswith("_share") or key in {"utilization", "availability", "reserve", "labor_reduction_share"}:
        return "%"
    return _param_unit(key) or "—"


def _reasons_text(items: Optional[list], limit: int = 3) -> str:
    if not items:
        return "—"
    cleaned = []
    for raw in items:
        text = str(raw).strip()
        if not text:
            continue
        # drop leading checkmarks noise for compactness
        text = text.lstrip("✓✗•-— ").strip()
        if len(text) > 140:
            text = text[:137] + "…"
        cleaned.append(f"• {text}")
        if len(cleaned) >= limit:
            break
    more = len(items) - len(cleaned)
    if more > 0:
        cleaned.append(f"• … ещё {more}")
    return "\n".join(cleaned) if cleaned else "—"


def _metric_ru(key: str, fallback: Optional[str] = None) -> str:
    return METRIC_LABELS_RU.get(key, fallback or key)


def _work_mode_ru(value: Optional[str]) -> str:
    if not value:
        return "Нет данных"
    return WORK_MODE_LABELS_RU.get(str(value), str(value))


def _status_ru(status: Optional[str]) -> str:
    if not status:
        return "Нет данных"
    return STATUS_LABELS_RU.get(status, status)


def _project_status_ru(status: Optional[str], is_demo: bool = False) -> str:
    if is_demo:
        return "Пример"
    if not status:
        return "Нет данных"
    return PROJECT_STATUS_RU.get(str(status), str(status))


def _strip_dataset_prefix(key: str) -> str:
    for prefix in (
        "ds_warehouse_",
        "ds_airport_",
        "ds_medical_",
        "ds_med_",
        "ds_",
    ):
        if key.startswith(prefix):
            return key[len(prefix) :]
    return key


def _param_label(key: str) -> str:
    from app.services.dataset_param_labels import DATASET_PARAM_LABELS_RU

    for candidate in (key, _strip_dataset_prefix(key)):
        if candidate in OBJECT_PARAM_META:
            return OBJECT_PARAM_META[candidate][0]
        if candidate in ASSUMPTION_LABELS_RU:
            return ASSUMPTION_LABELS_RU[candidate]
        if candidate in FORMULA_LABELS_RU:
            return FORMULA_LABELS_RU[candidate]
        if candidate in BREAKDOWN_LABELS_RU:
            return BREAKDOWN_LABELS_RU[candidate]
        if candidate in METRIC_LABELS_RU:
            return METRIC_LABELS_RU[candidate]
        if candidate in DATASET_PARAM_LABELS_RU:
            return DATASET_PARAM_LABELS_RU[candidate]
    # Не отдаём snake_case комиссии — нейтральная подпись
    base = _strip_dataset_prefix(key)
    if base != key:
        return f"Параметр датасета ({base.replace('_', ' ')})"
    if "_" in key and key.isascii():
        return f"Параметр ({key.replace('_', ' ')})"
    return key


def _param_unit(key: str) -> str:
    for candidate in (key, _strip_dataset_prefix(key)):
        if candidate in OBJECT_PARAM_META:
            return OBJECT_PARAM_META[candidate][1]
    unit_map = {
        "training_per_robot": "₽",
        "energy_per_robot_year": "₽/год",
        "connectivity_per_robot_year": "₽/год",
        "consumables_per_robot_year": "₽/год",
        "ops_staff_salary_year": "₽/год",
        "current_labor_cost_year": "₽/год",
        "additional_revenue": "₽/год",
        "raas_monthly_per_robot": "₽/мес",
        "unit_price": "₽",
        "software_cost": "₽",
        "maintenance_year": "₽/год",
        "implementation_cost": "₽",
        "years": "лет",
        "peak_demand": "оп/ч",
        "productivity": "оп/ч",
        "utilization": "%",
        "availability": "%",
        "reserve": "%",
    }
    return unit_map.get(key, "")


def _param_source(key: str) -> str:
    if key in ASSUMPTION_PARAM_KEYS and key not in {
        "peak_demand",
        "productivity",
        "robots_count",
        "current_labor_cost_year",
    }:
        if key in {
            "utilization",
            "availability",
            "reserve",
            "years",
            "infra_share_of_equipment",
            "integration_share",
            "commissioning_share",
            "training_per_robot",
            "contingency_share",
            "energy_per_robot_year",
            "connectivity_per_robot_year",
            "consumables_per_robot_year",
            "repair_share_of_price",
            "ops_staff_per_10_robots",
            "ops_staff_salary_year",
            "license_share_of_software",
            "labor_reduction_share",
            "additional_revenue",
            "raas_monthly_per_robot",
        }:
            return "Допущение модели"
    return "Параметры проекта"


def _assumption_label(code: str) -> str:
    return _param_label(code)


def _is_model_assumption(key: str) -> bool:
    if key in ASSUMPTION_LABELS_RU:
        return True
    # skip dataset dumps and technical prefixes
    if key.startswith("ds_") or key.startswith("raw_") or "." in key:
        return False
    return False


def _skip_object_param(key: str, value: Any) -> bool:
    if key.startswith("ds_") or key.startswith("_"):
        return True
    if isinstance(value, (dict, list)) and key not in {"required_processes", "cases"}:
        # nested blobs are not readable in a flat table
        if isinstance(value, dict):
            return True
    return False


def _count_statuses(solutions: list[dict]) -> dict[str, int]:
    counts = {"total": len(solutions), "suitable": 0, "needs_review": 0, "excluded": 0}
    for s in solutions:
        st = s.get("status")
        if st in counts:
            counts[st] += 1
    return counts


def _object_type_note(project: Any) -> str:
    for attr in ("object_type_name", "object_type_code"):
        val = getattr(project, attr, None)
        if val:
            return str(val)
    ot = getattr(project, "object_type", None)
    if ot is not None:
        name = getattr(ot, "name_ru", None) or getattr(ot, "name", None) or getattr(ot, "code", None)
        if name:
            return str(name)
    params = project.object_params or {}
    if isinstance(params, dict):
        for key in ("object_type", "facility_type", "zone_type"):
            if params.get(key):
                code = str(params[key])
                mapping = {
                    "warehouse": "Склад",
                    "airport": "Аэропорт",
                    "medical": "Медицинское учреждение",
                    "hospital": "Медицинское учреждение",
                }
                return mapping.get(code, code)
        # heuristic from description / name
    name = (getattr(project, "name", None) or "").lower()
    desc = (getattr(project, "description", None) or "").lower()
    blob = f"{name} {desc}"
    if "склад" in blob:
        return "Склад"
    if "аэропорт" in blob:
        return "Аэропорт"
    if "медицин" in blob or "мед" in blob:
        return "Медицинское учреждение"
    return "Не указан"


def _solution_sort_key(s: dict) -> tuple:
    status_rank = {"suitable": 0, "needs_review": 1, "excluded": 2}
    score = s.get("match_score")
    try:
        score_v = -float(score) if score is not None else 0.0
    except (TypeError, ValueError):
        score_v = 0.0
    selected = 0 if s.get("selected") else 1
    return (selected, status_rank.get(s.get("status"), 9), score_v)


def _top_solutions(solutions: list[dict], limit: int = 12) -> list[dict]:
    selected = [s for s in solutions if s.get("selected")]
    preferred = [s for s in solutions if s.get("status") in ("suitable", "needs_review")]
    pool = selected or preferred or list(solutions)
    return sorted(pool, key=_solution_sort_key)[:limit]


def _econ_value(economics: Optional[dict], path: str) -> Any:
    if not economics or not isinstance(economics, dict):
        return None
    if path == "robots_count":
        return economics.get("robots_count")
    if path == "capex":
        block = economics.get("capex")
        return block.get("capex") if isinstance(block, dict) else block
    if path == "opex_annual":
        block = economics.get("opex")
        return block.get("opex_annual") if isinstance(block, dict) else economics.get("opex_annual")
    if path == "annual_effect":
        block = economics.get("annual_effect")
        return block.get("annual_effect") if isinstance(block, dict) else block
    if path == "payback_years":
        block = economics.get("payback")
        return block.get("payback_years") if isinstance(block, dict) else economics.get("payback_years")
    if path == "roi_percent":
        block = economics.get("roi")
        return block.get("roi_percent") if isinstance(block, dict) else economics.get("roi_percent")
    if path == "tco":
        block = economics.get("tco")
        return block.get("tco") if isinstance(block, dict) else block
    return economics.get(path)


def _formulas_from_economics(economics: Optional[dict]) -> dict[str, Any]:
    if not economics or not isinstance(economics, dict):
        return {}
    formulas = economics.get("formulas")
    if isinstance(formulas, dict) and formulas:
        return formulas
    collected: dict[str, Any] = {}
    for key in ("capex", "opex", "annual_effect", "payback", "roi", "tco"):
        block = economics.get(key)
        if isinstance(block, dict) and block.get("formula"):
            collected[key] = block.get("formula")
    return collected


def _what_if_metric_getter(key: str) -> Callable[[dict], Any]:
    def getter(r: dict) -> Any:
        if not isinstance(r, dict):
            return None
        if key == "capex":
            c = r.get("capex")
            return c.get("capex") if isinstance(c, dict) else r.get("capex")
        if key == "opex_annual":
            o = r.get("opex")
            return o.get("opex_annual") if isinstance(o, dict) else r.get("opex_annual")
        if key == "annual_effect":
            e = r.get("annual_effect")
            return e.get("annual_effect") if isinstance(e, dict) else r.get("annual_effect")
        if key == "payback_years":
            p = r.get("payback")
            return p.get("payback_years") if isinstance(p, dict) else r.get("payback_years")
        if key == "roi_percent":
            roi = r.get("roi")
            return roi.get("roi_percent") if isinstance(roi, dict) else r.get("roi_percent")
        if key == "tco":
            t = r.get("tco")
            return t.get("tco") if isinstance(t, dict) else r.get("tco")
        return r.get(key)

    return getter


def _set_col_widths(ws, widths: dict[str, float]) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _style_table(
    ws,
    header_row: int = 1,
    wrap_cols: Optional[set[int]] = None,
    money_cols: Optional[set[int]] = None,
    percent_cols: Optional[set[int]] = None,
    auto_filter: bool = True,
) -> None:
    wrap_cols = wrap_cols or set()
    money_cols = money_cols or set()
    percent_cols = percent_cols or set()
    if ws.max_row < header_row or ws.max_column < 1:
        return
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        cell.border = THIN
    for r in range(header_row + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.border = THIN
            cell.alignment = Alignment(
                wrap_text=c in wrap_cols,
                vertical="top",
                horizontal="left",
            )
            if r % 2 == 0:
                cell.fill = ALT_ROW_FILL
            if c in money_cols and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = '#,##0'
            if c in percent_cols and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = '0.0'
    if auto_filter and ws.max_row >= header_row:
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
        )
    ws.freeze_panes = f"A{header_row + 1}"
    ws.row_dimensions[header_row].height = 22


def _append_kv_block(ws, rows: list[tuple[str, Any]], start_row: int = 1) -> int:
    r = start_row
    for label, value in rows:
        ws.cell(r, 1, label).font = Font(bold=True, color="334155")
        ws.cell(r, 2, value if value is not None else "—")
        ws.cell(r, 1).alignment = Alignment(vertical="top")
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    return r


def export_project_excel(
    project: Any,
    solutions: list[dict],
    economics: Optional[dict] = None,
    scenarios: Optional[list] = None,
    assumptions: Optional[dict] = None,
    sources: Optional[list] = None,
    what_if: Optional[dict] = None,
) -> bytes:
    is_demo = bool(getattr(project, "is_demo", False))
    counts = _count_statuses(solutions or [])
    today = datetime.now().strftime("%d.%m.%Y %H:%M")
    economics = economics if isinstance(economics, dict) else None
    object_type = _object_type_note(project)
    project_name = getattr(project, "name", None) or "Без названия"

    wb = Workbook()

    # 1. Титул — обложка отчёта
    ws_title = wb.active
    ws_title.title = "Титул"
    ws_title.merge_cells("A1:B1")
    ws_title["A1"] = PLATFORM_NAME
    ws_title["A1"].font = TITLE_FONT
    ws_title["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws_title.row_dimensions[1].height = 36

    ws_title.merge_cells("A2:B2")
    ws_title["A2"] = "Отчёт по проекту"
    ws_title["A2"].font = SUBTITLE_FONT

    ws_title.merge_cells("A3:B3")
    ws_title["A3"] = project_name
    ws_title["A3"].font = Font(bold=True, size=14)

    cover_rows = [
        ("Тип объекта", object_type),
        ("Дата формирования", today),
        ("Статус проекта", "Пример" if is_demo else "Рабочий проект"),
        ("Регион", _dash(getattr(project, "region", None))),
        ("Режим работы", _work_mode_ru(getattr(project, "work_mode", None))),
        ("", ""),
        ("Результаты подбора", ""),
        ("Всего рассмотрено решений", counts["total"]),
        ("Предварительно подходит", counts["suitable"]),
        ("Требует проверки", counts["needs_review"]),
        ("Не подходит", counts["excluded"]),
    ]
    if economics:
        cover_rows.extend(
            [
                ("", ""),
                ("Ключевые показатели экономики", ""),
                (
                    _metric_ru("capex"),
                    f"{_money_label(_econ_value(economics, 'capex'))} ₽",
                ),
                (
                    f"{_metric_ru('opex')} роботизации",
                    f"{_money_label(_econ_value(economics, 'opex_annual'))} ₽/год",
                ),
                (
                    f"{_metric_ru('tco')} роботизации",
                    f"{_money_label(_econ_value(economics, 'tco'))} ₽",
                ),
                (
                    "Полный выход (с оставшимся трудом)",
                    f"{_money_label((economics.get('tco') or {}).get('total_outlay_with_remaining_labor') if isinstance(economics.get('tco'), dict) else None)} ₽",
                ),
                (
                    _metric_ru("npv"),
                    f"{_money_label((economics.get('npv') or {}).get('npv') if isinstance(economics.get('npv'), dict) else None)} ₽",
                ),
                (
                    _metric_ru("roi"),
                    f"{_dash(_econ_value(economics, 'roi_percent'))} %",
                ),
                (
                    _metric_ru("payback"),
                    f"{_dash(_econ_value(economics, 'payback_years'))} лет",
                ),
            ]
        )
    cover_rows.extend(
        [
            ("", ""),
            (
                "Важно",
                (
                    "Результат является предварительной оценкой и требует верификации "
                    "при обследовании объекта."
                ),
            ),
            (
                "Примечание",
                (
                    "Значения рассчитаны платформой на основе параметров проекта и каталога решений."
                ),
            ),
        ]
    )
    _append_kv_block(ws_title, cover_rows, start_row=5)
    for r in range(5, ws_title.max_row + 1):
        label = ws_title.cell(r, 1).value
        if label in {"Результаты подбора", "Ключевые показатели экономики", "Важно"}:
            ws_title.cell(r, 1).fill = SECTION_FILL
            ws_title.cell(r, 2).fill = SECTION_FILL
            ws_title.cell(r, 1).font = Font(bold=True, color="1F4E79")
    _set_col_widths(ws_title, {"A": 34, "B": 62})

    # 2. Объект (проект + ключевые параметры)
    ws_proj = wb.create_sheet("Объект")
    ws_proj.append(["Параметр", "Значение"])
    for label, value in [
        ("Название", project_name),
        ("Тип объекта", object_type),
        ("Регион", _dash(getattr(project, "region", None))),
        ("Режим работы", _dash(getattr(project, "work_mode", None))),
        ("Статус", _project_status_ru(getattr(project, "status", None), is_demo=is_demo)),
        ("Описание", _dash(getattr(project, "description", None))),
        ("Тип данных", "Данные проекта"),
    ]:
        ws_proj.append([label, value])
    _style_table(ws_proj, wrap_cols={2}, auto_filter=False)
    _set_col_widths(ws_proj, {"A": 28, "B": 70})

    # 3. Параметры объекта (детализация)
    ws_params = wb.create_sheet("Параметры объекта")
    ws_params.append(["Параметр", "Значение", "Единица", "Источник"])
    params = project.object_params if isinstance(getattr(project, "object_params", None), dict) else {}
    written = 0
    if params:
        for key, value in params.items():
            if _skip_object_param(str(key), value):
                continue
            unit = _share_unit(str(key)) if str(key) in ASSUMPTION_LABELS_RU else (_param_unit(str(key)) or "—")
            display = _format_share_or_number(str(key), value) if isinstance(value, (int, float)) else _fmt_cell(value)
            ws_params.append(
                [
                    _param_label(str(key)),
                    display,
                    unit or "—",
                    _param_source(str(key)),
                ]
            )
            written += 1
    if written == 0:
        ws_params.append(["Нет данных в исходных материалах", "—", "—", "—"])
    _style_table(ws_params, wrap_cols={1, 2}, money_cols={2})
    _set_col_widths(ws_params, {"A": 40, "B": 28, "C": 12, "D": 22})

    # 4. Подобранные решения — сначала подходящие/проверка, короткие причины
    ws_sols = wb.create_sheet("Роботы")
    ws_sols.append(
        [
            "№",
            "Название",
            "Статус",
            "Оценка, %",
            "Краткий вывод",
            "Замечания / исключения",
            "Источник данных",
            "Файл",
        ]
    )
    ordered = sorted(solutions or [], key=_solution_sort_key)
    for idx, s in enumerate(ordered, start=1):
        origin = s.get("data_origin")
        status = s.get("status")
        if status == "suitable":
            summary = "Решение соответствует проверенным условиям"
        elif status == "needs_review":
            summary = "Есть неизвестные или неполные характеристики — требуется проверка"
        elif status == "excluded":
            summary = "Не соответствует обязательным условиям"
        else:
            summary = "Нет данных"
        reasons = s.get("match_reasons") or []
        exclusions = s.get("exclusion_reasons") or []
        # for suitable/needs_review show match reasons; for excluded show exclusions
        brief = _reasons_text(reasons if status != "excluded" else exclusions, limit=3)
        notes = _reasons_text(exclusions if status != "excluded" else reasons, limit=2)
        ws_sols.append(
            [
                idx,
                _dash(s.get("robot_name")),
                _status_ru(status),
                _dash(s.get("match_score")),
                f"{summary}\n{brief}" if brief != "—" else summary,
                notes,
                _origin_label(origin, is_demo=origin == "demo" or (is_demo and not origin)),
                _dash(s.get("source_file")),
            ]
        )
    if not ordered:
        ws_sols.append([1, "Нет подобранных решений", "—", "—", "Сначала выполните подбор", "—", "—", "—"])
    _style_table(ws_sols, wrap_cols={5, 6}, percent_cols={4})
    _set_col_widths(ws_sols, {"A": 6, "B": 28, "C": 18, "D": 11, "E": 48, "F": 36, "G": 22, "H": 26})
    for r in range(2, ws_sols.max_row + 1):
        ws_sols.row_dimensions[r].height = 48

    # 5. Сравнение — компактная матрица топ-5
    ws_cmp = wb.create_sheet("Сравнение")
    top = _top_solutions(solutions or [], limit=5)
    header = ["Показатель"] + [
        str(s.get("robot_name") or f"Решение {i + 1}")[:28] for i, s in enumerate(top)
    ]
    if not top:
        header = ["Показатель", "Нет данных"]
    ws_cmp.append(header)
    for field, label in COMPARISON_FIELDS:
        row: list[Any] = [label]
        if not top:
            row.append("Нет данных")
        else:
            for s in top:
                val = s.get(field)
                if field == "status":
                    val = _status_ru(val if isinstance(val, str) else None)
                elif field == "data_origin":
                    val = _origin_label(val if isinstance(val, str) else None)
                row.append(_dash(val))
        ws_cmp.append(row)
    _style_table(ws_cmp, wrap_cols=set(range(1, len(header) + 1)), auto_filter=False)
    widths = {"A": 24}
    for i in range(2, len(header) + 1):
        widths[get_column_letter(i)] = 18
    _set_col_widths(ws_cmp, widths)

    # 5b. Симуляция — метрики из последнего расчёта / fleet (без фиктивных значений)
    ws_sim = wb.create_sheet("Симуляция")
    ws_sim.append(["Показатель", "Значение", "Примечание"])
    sim_note = (
        "Дискретно-событийная модель по сохранённому плану; "
        "не полноценная физическая симуляция."
    )
    robots_n = _econ_value(economics, "robots_count") if economics else None
    ws_sim.append(["Количество роботов (активный сценарий)", _dash(robots_n), sim_note])
    fleet_rows = [
        sc
        for sc in (scenarios or [])
        if str(sc.get("code") or "").startswith("fleet_")
    ]
    if fleet_rows:
        ws_sim.append(["Сценарии парка", "", ""])
        ws_sim.append(
            [
                "Конфигурация",
                "Роботов",
                "Покрытие пика / загрузка / простой (оценка)",
            ]
        )
        for sc in fleet_rows:
            res = sc.get("results") if isinstance(sc.get("results"), dict) else {}
            ops = res.get("operations") if isinstance(res.get("operations"), dict) else {}
            m = _scenario_metrics(res)
            cov = ops.get("peak_coverage")
            util = ops.get("utilization_estimate")
            idle = ops.get("idle_estimate")
            detail = (
                f"покрытие={_dash(cov)}; загрузка={_dash(util)}; простой={_dash(idle)}"
                + ("; оценочно" if ops.get("estimated") else "")
            )
            if ops.get("note"):
                detail = f"{detail}; {ops.get('note')}"
            ws_sim.append(
                [
                    _dash(sc.get("name_ru") or sc.get("code")),
                    _dash(m.get("robots_count") or res.get("robots_count")),
                    detail or "Нет данных симуляции",
                ]
            )
    else:
        ws_sim.append(
            [
                "Сценарии парка",
                "Нет данных",
                "Запустите сравнение конфигураций (4/6/8) в интерфейсе",
            ]
        )
    ws_sim.append(
        [
            "Ограничение модели",
            "—",
            "Метрики производительности и узких мест берутся из модели маршрутов/спроса; "
            "при отсутствии данных ячейка содержит «Нет данных», а не подстановку.",
        ]
    )
    _style_table(ws_sim, wrap_cols={3}, auto_filter=False)
    _set_col_widths(ws_sim, {"A": 36, "B": 14, "C": 64})

    # 6. Экономика
    ws_econ = wb.create_sheet("Экономика")
    ws_econ.append(["Показатель", "Значение", "Единица", "Описание"])
    scenario_name = None
    if economics:
        scenario_name = economics.get("scenario_name_ru") or economics.get("scenario")
        code_map = {"purchase": "Покупка оборудования", "raas": "Роботы как услуга", "baseline": "Без роботизации"}
        if isinstance(scenario_name, str) and scenario_name in code_map:
            scenario_name = code_map[scenario_name]
    ws_econ.append(["Сценарий", _dash(scenario_name), "—", "Активный сценарий расчёта"])
    for key, title, unit in [
        ("robots_count", "Количество роботов", "шт"),
        ("capex", _metric_ru("capex"), "₽"),
        ("opex_annual", f"{_metric_ru('opex')} роботизации", "₽/год"),
        ("annual_effect", _metric_ru("annual_effect"), "₽/год"),
        ("tco", f"{_metric_ru('tco')} роботизации", "₽"),
        ("roi_percent", _metric_ru("roi"), "%"),
        ("payback_years", _metric_ru("payback"), "лет"),
    ]:
        raw = _econ_value(economics, key)
        value = _format_money(raw) if unit in {"₽", "₽/год"} else _dash(raw)
        ws_econ.append([title, value, unit, ECONOMICS_DESCRIPTIONS.get(key, "")])
    if economics:
        tco_block = economics.get("tco") if isinstance(economics.get("tco"), dict) else {}
        outlay = tco_block.get("total_outlay_with_remaining_labor")
        rem = tco_block.get("remaining_manual_labor_annual")
        if rem is not None:
            ws_econ.append(
                [
                    "Оставшийся ручной труд",
                    _format_money(rem),
                    "₽/год",
                    "Не входит в совокупную стоимость владения",
                ]
            )
        if outlay is not None:
            ws_econ.append(
                [
                    "Полный денежный выход",
                    _format_money(outlay),
                    "₽",
                    "Совокупная стоимость + оставшийся труд × горизонт",
                ]
            )
        years = (economics.get("inputs") or {}).get("years") if isinstance(economics.get("inputs"), dict) else tco_block.get("years")
        discount = None
        if isinstance(economics.get("npv"), dict):
            discount = economics["npv"].get("discount_rate")
        if years is not None:
            ws_econ.append(
                [
                    "Горизонт анализа",
                    _dash(years),
                    "лет",
                    "Годы для совокупной стоимости / ЧПС / денежного потока",
                ]
            )
        if discount is not None:
            ws_econ.append(
                [
                    "Ставка дисконтирования",
                    f"{float(discount) * 100:.1f}",
                    "%",
                    "Допущение модели для чистой приведённой стоимости",
                ]
            )
        npv_v = economics.get("npv", {}).get("npv") if isinstance(economics.get("npv"), dict) else None
        irr_v = economics.get("irr", {}).get("irr") if isinstance(economics.get("irr"), dict) else None
        ws_econ.append(
            [
                _metric_ru("npv"),
                _format_money(npv_v) if npv_v is not None else "Нет данных",
                "₽",
                "Дисконтированные денежные потоки",
            ]
        )
        ws_econ.append(
            [
                _metric_ru("irr"),
                _dash(irr_v),
                "%",
                "При невозможности расчёта — «Нет данных»",
            ]
        )
        dep = economics.get("depreciation") if isinstance(economics.get("depreciation"), dict) else {}
        if dep.get("applicable") and dep.get("annual_depreciation") is not None:
            ws_econ.append(["Амортизация капитальных затрат (бухгалтерская)", "", "", ""])
            method_raw = dep.get("method_label_ru") or dep.get("method") or "straight_line"
            method_label = METRIC_LABELS_RU.get(str(method_raw), method_raw)
            ws_econ.append(
                [
                    "Метод",
                    method_label,
                    "",
                    "Допущение команды; не денежный отток",
                ]
            )
            ws_econ.append(
                [
                    "Срок службы",
                    _dash(dep.get("useful_life_years")),
                    "лет",
                    "Допущение (в ТЗ число не задано)",
                ]
            )
            ws_econ.append(
                [
                    "База",
                    _format_money(dep.get("depreciable_base")),
                    "₽",
                    str(dep.get("asset_label") or _metric_ru("capex")),
                ]
            )
            ws_econ.append(
                [
                    "Амортизация / год",
                    _format_money(dep.get("annual_depreciation")),
                    "₽/год",
                    "Не входит в денежный поток / ЧПС / ВНД",
                ]
            )
            effect = economics.get("annual_effect") if isinstance(economics.get("annual_effect"), dict) else {}
            if effect.get("with_depreciation") is not None:
                ws_econ.append(
                    [
                        "Годовой эффект с амортизацией",
                        _format_money(effect.get("with_depreciation")),
                        "₽/год",
                        "Денежный эффект − амортизация (бухгалтерский вид)",
                    ]
                )
        cfs = economics.get("cashflows") if isinstance(economics.get("cashflows"), list) else []
        if cfs:
            ws_econ.append(["Денежные потоки", "", "", ""])
            ws_econ.append(
                ["Год", "Денежный поток (CF)", "Накоплено", "Амортизация (бух., не в CF)"]
            )
            for row in cfs:
                if isinstance(row, dict):
                    ws_econ.append(
                        [
                            row.get("year"),
                            _format_money(row.get("cashflow")),
                            _format_money(row.get("cumulative")),
                            (
                                _format_money(row.get("depreciation_accounting"))
                                if row.get("depreciation_accounting") is not None
                                else ""
                            ),
                        ]
                    )

    # section: CAPEX breakdown
    if economics:
        capex_block = economics.get("capex") if isinstance(economics.get("capex"), dict) else {}
        breakdown = capex_block.get("breakdown") or {}
        if breakdown:
            ws_econ.append(["Состав капитальных затрат", "", "", ""])
            sources = capex_block.get("component_sources") or {}
            for k, v in breakdown.items():
                if k == "ops_staff_count":
                    continue
                src = sources.get(str(k)) if isinstance(sources, dict) else None
                origin = (src or {}).get("origin") if isinstance(src, dict) else None
                note = (src or {}).get("note") if isinstance(src, dict) else None
                desc = "Статья капитальных затрат"
                if origin == "model_assumption":
                    desc = f"Допущение модели: {note or 'доля/норматив'}"
                elif note:
                    desc = str(note)
                ws_econ.append(
                    [
                        BREAKDOWN_LABELS_RU.get(str(k), str(k)),
                        _format_money(v),
                        "₽",
                        desc,
                    ]
                )
            if capex_block.get("breakdown_matches_total") is False:
                ws_econ.append(
                    [
                        "⚠ Сумма статей ≠ капитальным затратам",
                        _format_money(capex_block.get("breakdown_sum")),
                        "₽",
                        "Внутренняя проверка не пройдена",
                    ]
                )
        opex_block = economics.get("opex") if isinstance(economics.get("opex"), dict) else {}
        ob = opex_block.get("breakdown") or {}
        if ob:
            ws_econ.append(["Состав эксплуатационных расходов", "", "", ""])
            for k, v in ob.items():
                unit = "чел" if k == "ops_staff_count" else "₽/год"
                ws_econ.append(
                    [
                        BREAKDOWN_LABELS_RU.get(str(k), str(k)),
                        _format_money(v) if unit != "чел" else _dash(v),
                        unit,
                        "Статья эксплуатационных затрат",
                    ]
                )
    _style_table(ws_econ, wrap_cols={4}, money_cols={2}, auto_filter=False)
    for r in range(2, ws_econ.max_row + 1):
        if ws_econ.cell(r, 1).value in {"Состав капитальных затрат", "Состав эксплуатационных расходов"}:
            for c in range(1, 5):
                ws_econ.cell(r, c).fill = SECTION_FILL
                ws_econ.cell(r, c).font = Font(bold=True, color="1F4E79")
    _set_col_widths(ws_econ, {"A": 34, "B": 16, "C": 10, "D": 48})

    # 7. Как рассчитано
    ws_how = wb.create_sheet("Как рассчитано")
    ws_how.append(["Показатель", "Формула / пояснение", "Источник"])
    formulas = _formulas_from_economics(economics)
    if formulas:
        for k, v in formulas.items():
            label = FORMULA_LABELS_RU.get(str(k), _assumption_label(str(k)) if str(k) in ASSUMPTION_LABELS_RU else str(k))
            # avoid raw english snake_case headers when possible
            if label == str(k) and "_" in label:
                label = label.replace("_", " ")
            ws_how.append([label, str(v), "Рассчитано платформой"])
    else:
        ws_how.append(
            [
                "Примечание",
                "Показатели рассчитаны платформой по встроенным формулам экономики.",
                "Рассчитано платформой",
            ]
        )
    _style_table(ws_how, wrap_cols={2}, auto_filter=False)
    _set_col_widths(ws_how, {"A": 24, "B": 70, "C": 22})

    # 8. Сценарии
    ws_sc = wb.create_sheet("Сценарии")
    ws_sc.append(
        [
            "Сценарий",
            "Роботов",
            "Капитальные затраты, ₽",
            "Эксплуатационные расходы, ₽/год",
            "Годовой эффект, ₽",
            "Окупаемость, лет",
            "Рентабельность, %",
            "Совокупная стоимость владения, ₽",
            "Чистая приведённая стоимость, ₽",
        ]
    )
    for sc in scenarios or []:
        m = _scenario_metrics(sc.get("results") or {})
        ws_sc.append(
            [
                _dash(sc.get("name_ru") or sc.get("code")),
                _dash(m.get("robots_count")),
                _format_money(m.get("capex")),
                _format_money(m.get("opex_annual")),
                _format_money(m.get("annual_effect")),
                _dash(m.get("payback_years")),
                _dash(m.get("roi_percent")),
                _format_money(m.get("tco")),
                _format_money(m.get("npv")),
            ]
        )
    if not scenarios:
        ws_sc.append(["Нет сохранённых сценариев", "—", "—", "—", "—", "—", "—", "—", "—"])
    _style_table(ws_sc, money_cols={3, 4, 5, 8, 9}, percent_cols={7})
    _set_col_widths(ws_sc, {"A": 26, "B": 10, "C": 14, "D": 14, "E": 16, "F": 14, "G": 10, "H": 14, "I": 14})

    # 9. Что если
    ws_wi = wb.create_sheet("Что если")
    ws_wi.append(["Показатель", "Исходное", "Изменённое", "Изменение", "Изменение, %"])
    if what_if and isinstance(what_if, dict):
        base = what_if.get("base") or {}
        alt = what_if.get("alternative") or {}
        delta = what_if.get("delta") or {}
        metrics = [
            ("Капитальные затраты, ₽", "capex", True),
            ("Эксплуатационные расходы, ₽/год", "opex_annual", True),
            ("Годовой эффект, ₽", "annual_effect", True),
            ("Количество роботов", "robots_count", False),
            ("Срок окупаемости, лет", "payback_years", False),
            ("Рентабельность, %", "roi_percent", False),
            ("Совокупная стоимость владения, ₽", "tco", True),
        ]
        for title, key, money in metrics:
            getter = _what_if_metric_getter(key)
            b = getter(base) if isinstance(base, dict) else None
            a = getter(alt) if isinstance(alt, dict) else None
            d = delta.get(key) if isinstance(delta, dict) else None
            if d is None and b is not None and a is not None:
                try:
                    d = float(a) - float(b)
                except (TypeError, ValueError):
                    d = None
            pct = None
            try:
                if b not in (None, 0) and d is not None:
                    pct = round(100.0 * float(d) / float(b), 1)
            except (TypeError, ValueError, ZeroDivisionError):
                pct = None
            ws_wi.append(
                [
                    title,
                    _format_money(b) if money else _dash(b),
                    _format_money(a) if money else _dash(a),
                    _format_money(d) if money else _dash(d),
                    _dash(pct),
                ]
            )
        changes = what_if.get("changes") or {}
        if changes:
            ws_wi.append(["Изменённые допущения", "", "", "", ""])
            for k, v in changes.items():
                ws_wi.append(
                    [
                        _assumption_label(str(k)),
                        "—",
                        _format_share_or_number(str(k), v),
                        "—",
                        "—",
                    ]
                )
    else:
        ws_wi.append(
            [
                "Анализ «что если» ещё не сохранялся",
                "—",
                "—",
                "—",
                "Измените допущения на шаге «Сценарии и что если» и скачайте отчёт снова",
            ]
        )
    _style_table(ws_wi, money_cols={2, 3, 4}, percent_cols={5}, auto_filter=False)
    _set_col_widths(ws_wi, {"A": 30, "B": 14, "C": 14, "D": 14, "E": 14})

    # 10. Допущения — только модельные, без ds_*
    ws_ass = wb.create_sheet("Допущения")
    ws_ass.append(["Показатель", "Значение", "Единица", "Пояснение"])
    assum = assumptions or (economics or {}).get("assumptions") if economics else assumptions
    assum = assum or {}
    wrote_ass = 0
    if isinstance(assum, dict) and assum:
        for k, v in assum.items():
            key = str(k)
            if not _is_model_assumption(key):
                continue
            unit = _share_unit(key)
            ws_ass.append(
                [
                    _assumption_label(key),
                    _format_share_or_number(key, v),
                    unit or "—",
                    "Допущение модели расчёта",
                ]
            )
            wrote_ass += 1
    if wrote_ass == 0:
        ws_ass.append(["Нет сохранённых допущений", "—", "—", "—"])
    _style_table(ws_ass, wrap_cols={1}, money_cols={2})
    _set_col_widths(ws_ass, {"A": 42, "B": 14, "C": 10, "D": 30})

    # 11. Источники и допущения (сводка + ссылки)
    ws_src = wb.create_sheet("Источники и допущения")
    ws_src.append(["Раздел", "Наименование", "Значение / ссылка", "Комментарий"])
    if isinstance(assum, dict) and assum:
        for k, v in assum.items():
            key = str(k)
            if not _is_model_assumption(key):
                continue
            ws_src.append(
                [
                    "Допущение",
                    _assumption_label(key),
                    _format_share_or_number(key, v),
                    "Допущение модели расчёта",
                ]
            )
    for src in sources or []:
        if isinstance(src, dict):
            ws_src.append(
                [
                    "Источник",
                    _dash(src.get("name")),
                    _dash(src.get("url")),
                    _dash(src.get("description")),
                ]
            )
        else:
            ws_src.append(
                [
                    "Источник",
                    _dash(getattr(src, "name", None)),
                    _dash(getattr(src, "url", None)),
                    _dash(getattr(src, "description", None)),
                ]
            )
    seen_files: set[str] = set()
    for s in solutions or []:
        sf = s.get("source_file")
        if sf and sf not in seen_files:
            seen_files.add(sf)
            row_no = s.get("source_row")
            ws_src.append(
                [
                    "Каталог",
                    sf,
                    f"Строка {row_no}" if row_no is not None else "—",
                    _origin_label(s.get("data_origin")),
                ]
            )
    if is_demo:
        ws_src.append(
            [
                "Источник",
                "Каталог платформы",
                "—",
                "Данные проекта",
            ]
        )
    if ws_src.max_row == 1:
        ws_src.append(["—", "Нет данных", "—", "—"])
    _style_table(ws_src, wrap_cols={4})
    _set_col_widths(ws_src, {"A": 14, "B": 40, "C": 22, "D": 48})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_project_pdf(
    project: Any,
    economics: Optional[dict] = None,
    solutions_summary: Optional[list] = None,
    scenarios: Optional[list] = None,
    assumptions: Optional[dict] = None,
    sources: Optional[list] = None,
    what_if: Optional[dict] = None,
) -> bytes:
    font_name = _register_cyrillic_font()
    project_name = getattr(project, "name", None) or "Без названия"
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=f"Отчёт — {project_name}"[:180],
        author="Платформа подбора робототехники",
        creator="NETCORE",
        subject="Отчёт по проекту роботизации",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "RuTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=22,
        textColor=_PDF_INK,
        spaceAfter=4,
    )
    brand_style = ParagraphStyle(
        "RuBrand",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9,
        leading=12,
        textColor=_PDF_ORANGE,
        spaceAfter=2,
    )
    h2_style = ParagraphStyle(
        "RuH2",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=12,
        leading=16,
        textColor=_PDF_INK,
        spaceBefore=10,
        spaceAfter=4,
    )
    normal_style = ParagraphStyle(
        "RuNormal",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=10,
        leading=13,
        textColor=_PDF_INK,
    )
    small_style = ParagraphStyle(
        "RuSmall",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8,
        leading=11,
        textColor=_PDF_INK,
    )
    kpi_label_style = ParagraphStyle(
        "RuKpiLabel",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=7.5,
        leading=9,
        textColor=_PDF_MUTED,
    )
    kpi_value_style = ParagraphStyle(
        "RuKpiValue",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=11,
        leading=14,
        textColor=_PDF_INK,
    )
    footer_style = ParagraphStyle(
        "RuFooter",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8,
        leading=11,
        textColor=_PDF_MUTED,
    )

    def esc(text: Any) -> str:
        return (
            str(text)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    def p(text: Any, style: ParagraphStyle = normal_style) -> Paragraph:
        return Paragraph(esc(text), style)

    header_cell_style = ParagraphStyle(
        "RuHeaderCell",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8,
        leading=11,
        textColor=colors.white,
    )

    def make_table(data: list[list[Any]], col_widths: Optional[list] = None) -> Table:
        # Wrap long cells as Paragraphs for Cyrillic wrapping
        wrapped: list[list[Any]] = []
        for r_idx, row in enumerate(data):
            new_row = []
            for cell in row:
                if isinstance(cell, Paragraph):
                    new_row.append(cell)
                else:
                    style = header_cell_style if r_idx == 0 else small_style
                    new_row.append(Paragraph(esc(_dash(cell)), style))
            wrapped.append(new_row)
        table = Table(wrapped, colWidths=col_widths, hAlign="LEFT")
        style_cmds: list = [
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), _PDF_HEADER),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.Color(0.75, 0.78, 0.82)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for i in range(1, len(wrapped)):
            if i % 2 == 0:
                style_cmds.append(("BACKGROUND", (0, i), (-1, i), _PDF_ROW_ALT))
        table.setStyle(TableStyle(style_cmds))
        return table

    def _kpi_card(label: str, value: Any) -> Table:
        card = Table(
            [
                [Paragraph(esc(label), kpi_label_style)],
                [Paragraph(f"<b>{esc(_dash(value))}</b>", kpi_value_style)],
            ],
            colWidths=[78 * mm],
        )
        card.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), _PDF_CARD),
                    ("BOX", (0, 0), (-1, -1), 1.2, _PDF_CARD_LINE),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        return card

    def kpi_strip(cards: list[tuple[str, Any]]) -> Table:
        """Compact visual KPI grid (label + value cards), not a data dump."""
        grid_rows: list[list[Any]] = []
        for i in range(0, len(cards), 2):
            left = _kpi_card(*cards[i])
            right = _kpi_card(*cards[i + 1]) if i + 1 < len(cards) else Spacer(1, 1)
            row_t = Table([[left, right]], colWidths=[80 * mm, 80 * mm])
            row_t.setStyle(
                TableStyle(
                    [
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                        ("TOPPADDING", (0, 0), (-1, -1), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            grid_rows.append([row_t])
        wrap = Table(grid_rows, colWidths=[160 * mm])
        wrap.setStyle(
            TableStyle(
                [
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
        return wrap

    def section(title: str, flowables: list) -> KeepTogether:
        bar = HRFlowable(
            width="100%",
            thickness=2.2,
            color=_PDF_ORANGE,
            spaceBefore=2,
            spaceAfter=6,
            hAlign="LEFT",
        )
        return KeepTogether([p(title, h2_style), bar, *flowables, Spacer(1, 6)])

    is_demo = bool(getattr(project, "is_demo", False))
    economics = economics if isinstance(economics, dict) else None
    solutions = solutions_summary or []
    counts = _count_statuses(solutions)
    today = datetime.now().strftime("%d.%m.%Y")

    if what_if is None and economics and isinstance(economics.get("what_if"), dict):
        what_if = economics.get("what_if")
    if assumptions is None and economics:
        assumptions = economics.get("assumptions")

    story: list = []

    # Title page
    story.append(p("NETCORE · промышленный командный центр", brand_style))
    story.append(p(PLATFORM_NAME, title_style))
    story.append(
        HRFlowable(
            width="100%",
            thickness=3,
            color=_PDF_ORANGE,
            spaceBefore=2,
            spaceAfter=8,
            hAlign="LEFT",
        )
    )
    story.append(p(f"Проект: {getattr(project, 'name', None) or 'Нет данных'}", normal_style))
    story.append(p(f"Дата: {today}", normal_style))
    story.append(
        p(
            "Отчёт по проекту",
            normal_style,
        )
    )
    story.append(Spacer(1, 8))

    # Visual KPI strip (same numbers as before — presentation only)
    npv_val = (
        (economics or {}).get("npv", {}).get("npv")
        if isinstance((economics or {}).get("npv"), dict)
        else None
    )
    irr_val = (
        (economics or {}).get("irr", {}).get("irr")
        if isinstance((economics or {}).get("irr"), dict)
        else None
    )
    story.append(
        section(
            "Ключевые показатели",
            [
                kpi_strip(
                    [
                        (_metric_ru("capex"), _econ_value(economics, "capex")),
                        (
                            f"{_metric_ru('opex')} / год",
                            _econ_value(economics, "opex_annual"),
                        ),
                        (
                            f"{_metric_ru('roi')}, %",
                            _econ_value(economics, "roi_percent"),
                        ),
                        (
                            f"{_metric_ru('payback')}, лет",
                            _econ_value(economics, "payback_years"),
                        ),
                        (_metric_ru("npv"), npv_val),
                        (f"{_metric_ru('irr')}, %", irr_val),
                        ("Количество роботов", _econ_value(economics, "robots_count")),
                        ("Подходящих решений", counts["suitable"]),
                    ]
                )
            ],
        )
    )

    # Резюме — детальная таблица
    resume_rows = [
        ["Показатель", "Значение"],
        ["Всего решений", counts["total"]],
        ["Предварительно подходит", counts["suitable"]],
        ["Требует проверки", counts["needs_review"]],
        ["Не подходит", counts["excluded"]],
        [_metric_ru("capex"), _dash(_econ_value(economics, "capex"))],
        [
            f"{_metric_ru('opex')} роботизации",
            _dash(_econ_value(economics, "opex_annual")),
        ],
        [
            f"{_metric_ru('tco')} роботизации (без оставшегося ручного труда)",
            _dash(_econ_value(economics, "tco")),
        ],
        [
            "Полный денежный выход (Совокупная стоимость + оставшийся труд × горизонт)",
            _dash(
                (economics or {}).get("tco", {}).get("total_outlay_with_remaining_labor")
                if isinstance((economics or {}).get("tco"), dict)
                else None
            ),
        ],
        [
            f"{_metric_ru('roi')}, %",
            _dash(_econ_value(economics, "roi_percent")),
        ],
        [
            f"{_metric_ru('payback')}, лет",
            _dash(_econ_value(economics, "payback_years")),
        ],
        [_metric_ru("npv"), _dash(npv_val)],
        [f"{_metric_ru('irr')}, %", _dash(irr_val)],
        ["Количество роботов", _dash(_econ_value(economics, "robots_count"))],
    ]
    story.append(
        section(
            "Резюме",
            [make_table(resume_rows, col_widths=[90 * mm, 70 * mm])],
        )
    )

    # Параметры объекта
    params = project.object_params if isinstance(getattr(project, "object_params", None), dict) else {}
    param_data = [["Параметр", "Значение", "Единица", "Источник"]]
    wrote_params = 0
    if params:
        for key, value in params.items():
            if _skip_object_param(str(key), value):
                continue
            param_data.append(
                [
                    _param_label(str(key)),
                    _fmt_cell(value),
                    _param_unit(str(key)) or "—",
                    _param_source(str(key)),
                ]
            )
            wrote_params += 1
            if wrote_params >= 40:
                break
    if wrote_params == 0:
        param_data.append(["Нет данных", "Нет данных", "—", "—"])
    story.append(
        section(
            "Параметры объекта",
            [make_table(param_data, col_widths=[50 * mm, 50 * mm, 25 * mm, 35 * mm])],
        )
    )

    # Подобранные решения (top 12)
    top = _top_solutions(solutions, limit=12)
    sol_data = [["№", "Название", "Статус", "Оценка", "Происхождение"]]
    if top:
        for s in top:
            sol_data.append(
                [
                    _dash(s.get("robot_id")),
                    _dash(s.get("robot_name")),
                    _status_ru(s.get("status")),
                    _dash(s.get("match_score")),
                    _origin_label(s.get("data_origin")),
                ]
            )
    else:
        sol_data.append(["Нет данных", "—", "—", "—", "—"])
    story.append(
        section(
            "Подобранные решения",
            [make_table(sol_data, col_widths=[18 * mm, 52 * mm, 35 * mm, 18 * mm, 37 * mm])],
        )
    )

    # Экономика
    econ_data = [["Показатель", "Значение", "Единица", "Описание"]]
    scenario_name = None
    if economics:
        scenario_name = economics.get("scenario_name_ru") or economics.get("scenario")
        if isinstance(scenario_name, str):
            scenario_name = {
                "baseline": "Без роботизации",
                "purchase": "Покупка оборудования",
                "raas": "Аренда",
            }.get(scenario_name, scenario_name)
    econ_data.append(["Сценарий", _dash(scenario_name), "—", "Активный сценарий"])
    for key, title, unit, desc_key in [
        ("robots_count", "Количество роботов", "шт", "robots_count"),
        ("capex", "Капитальные затраты", "₽", "capex"),
        ("opex_annual", "Эксплуатационные затраты", "₽/год", "opex_annual"),
        ("annual_effect", "Годовой эффект", "₽/год", "annual_effect"),
        ("tco", "Совокупная стоимость владения", "₽", "tco"),
        ("roi_percent", "Рентабельность инвестиций", "%", "roi_percent"),
        ("payback_years", "Срок окупаемости", "лет", "payback_years"),
    ]:
        econ_data.append(
            [
                title,
                _dash(_econ_value(economics, key)),
                unit,
                ECONOMICS_DESCRIPTIONS.get(desc_key, ""),
            ]
        )
    dep = economics.get("depreciation") if isinstance(economics.get("depreciation"), dict) else {}
    if dep.get("applicable") and dep.get("annual_depreciation") is not None:
        econ_data.append(
            [
                "Амортизация капитальных затрат / год",
                _dash(dep.get("annual_depreciation")),
                "₽/год",
                "Бухгалтерский расход; не денежный отток",
            ]
        )
        econ_data.append(
            [
                "Срок амортизации",
                _dash(dep.get("useful_life_years")),
                "лет",
                "Допущение команды (в ТЗ число не задано)",
            ]
        )
        effect = economics.get("annual_effect") if isinstance(economics.get("annual_effect"), dict) else {}
        if effect.get("with_depreciation") is not None:
            econ_data.append(
                [
                    "Годовой эффект с амортизацией",
                    _dash(effect.get("with_depreciation")),
                    "₽/год",
                    "Денежный эффект − амортизация; не используется в ЧПС/ВНД",
                ]
            )
    story.append(
        section(
            "Экономика",
            [make_table(econ_data, col_widths=[55 * mm, 30 * mm, 20 * mm, 55 * mm])],
        )
    )

    # Как рассчитано
    formulas = _formulas_from_economics(economics)
    how_items: list = []
    if formulas:
        how_data = [["Показатель", "Формула"]]
        for k, v in formulas.items():
            label = FORMULA_LABELS_RU.get(str(k), _assumption_label(str(k)))
            how_data.append([label, str(v)])
        how_items.append(make_table(how_data, col_widths=[50 * mm, 110 * mm]))
    else:
        how_items.append(p("Рассчитано платформой", normal_style))
    story.append(section("Как рассчитано", how_items))

    # Денежные потоки
    cfs = (economics or {}).get("cashflows") if isinstance(economics, dict) else None
    if isinstance(cfs, list) and cfs:
        cf_data = [["Год", "Денежный поток", "Накоплено", "Дисконтированный"]]
        for row in cfs:
            if not isinstance(row, dict):
                continue
            cf_data.append(
                [
                    _dash(row.get("year")),
                    _dash(row.get("cashflow")),
                    _dash(row.get("cumulative")),
                    _dash(row.get("discounted")),
                ]
            )
        story.append(
            section(
                "Денежные потоки",
                [make_table(cf_data, col_widths=[25 * mm, 45 * mm, 45 * mm, 45 * mm])],
            )
        )
    else:
        story.append(
            section(
                "Денежные потоки",
                [p("Нет данных: выполните расчёт экономики с горизонтом планирования.", normal_style)],
            )
        )

    # Сценарии — всегда в отчёте (пустая таблица с пояснением, если данных нет)
    sc_data = [
        [
            "Сценарий",
            "Роботы",
            "Капитальные затраты, ₽",
            "Эксплуатационные расходы, ₽/год",
            "Годовой эффект, ₽",
            "Рентабельность, %",
            "Окупаемость, лет",
            "Чистая приведённая стоимость, ₽",
        ]
    ]
    code_map = {
        "baseline": "Без роботизации",
        "purchase": "Покупка оборудования",
        "raas": "Роботы как услуга",
    }
    wrote_sc = 0
    for sc in scenarios or []:
        m = _scenario_metrics(sc.get("results") or {})
        if not m:
            continue
        code = sc.get("code")
        name = sc.get("name_ru") or code_map.get(str(code or ""), code)
        sc_data.append(
            [
                _dash(name),
                _dash(m.get("robots_count")),
                _dash(m.get("capex")),
                _dash(m.get("opex_annual")),
                _dash(m.get("annual_effect")),
                _dash(m.get("roi_percent")),
                _dash(m.get("payback_years")),
                _dash(m.get("npv")),
            ]
        )
        wrote_sc += 1
    if wrote_sc == 0:
        sc_data.append(["Нет данных по сценариям", "—", "—", "—", "—", "—", "—", "—"])
    story.append(
        section(
            "Сравнение сценариев",
            [
                make_table(
                    sc_data,
                    col_widths=[
                        32 * mm,
                        14 * mm,
                        22 * mm,
                        22 * mm,
                        22 * mm,
                        14 * mm,
                        14 * mm,
                        20 * mm,
                    ],
                ),
                p(
                    "Сравнение сценариев «без роботизации / покупка / аренда» по текущим параметрам проекта.",
                    small_style,
                ),
            ],
        )
    )

    # What-if
    wi_items: list = []
    if what_if and isinstance(what_if, dict):
        wi_data = [["Показатель", "База", "Новое", "Δ", "Δ %"]]
        base = what_if.get("base") or {}
        alt = what_if.get("alternative") or {}
        delta = what_if.get("delta") or {}
        for title, key in [
            ("Капитальные затраты", "capex"),
            ("Эксплуатационные затраты, год", "opex_annual"),
            ("Годовой эффект", "annual_effect"),
            ("Роботов", "robots_count"),
            ("Срок окупаемости, лет", "payback_years"),
            ("Рентабельность, %", "roi_percent"),
            ("Совокупная стоимость владения, ₽", "tco"),
        ]:
            getter = _what_if_metric_getter(key)
            b = getter(base) if isinstance(base, dict) else None
            a = getter(alt) if isinstance(alt, dict) else None
            d = delta.get(key) if isinstance(delta, dict) else None
            if d is None and b is not None and a is not None:
                try:
                    d = float(a) - float(b)
                except (TypeError, ValueError):
                    d = None
            pct = None
            try:
                if b not in (None, 0) and d is not None:
                    pct = round(100.0 * float(d) / float(b), 2)
            except (TypeError, ValueError, ZeroDivisionError):
                pct = None
            wi_data.append([title, _dash(b), _dash(a), _dash(d), _dash(pct)])
        wi_items.append(make_table(wi_data, col_widths=[55 * mm, 25 * mm, 25 * mm, 25 * mm, 25 * mm]))
        changes = what_if.get("changes") or {}
        if isinstance(changes, dict) and changes:
            ch_data = [["Допущение", "Значение"]]
            for k, v in changes.items():
                ch_data.append([_assumption_label(str(k)), _format_share_or_number(str(k), v)])
            wi_items.append(Spacer(1, 4))
            wi_items.append(p("Изменённые допущения", normal_style))
            wi_items.append(make_table(ch_data, col_widths=[80 * mm, 80 * mm]))
    else:
        wi_items.append(
            p(
                "Анализ «что если» ещё не сохранялся. Откройте «Сценарии и что если», "
                "измените допущения — затем снова скачайте отчёт.",
                normal_style,
            )
        )
    story.append(section("Анализ «что если»", wi_items))

    # Допущения — только модельные (без ds_* и сырых ключей датасета)
    assum = assumptions or {}
    ass_data = [["Показатель", "Значение"]]
    wrote_ass = 0
    if isinstance(assum, dict) and assum:
        for k, v in assum.items():
            key = str(k)
            if not _is_model_assumption(key):
                continue
            ass_data.append([_assumption_label(key), _fmt_cell(v)])
            wrote_ass += 1
            if wrote_ass >= 40:
                break
    if wrote_ass == 0:
        ass_data.append(["Нет данных", "Нет данных"])
    story.append(
        section(
            "Допущения",
            [make_table(ass_data, col_widths=[100 * mm, 60 * mm])],
        )
    )

    # Источники
    src_data = [["Название", "Ссылка", "Тип"]]
    any_src = False
    for src in sources or []:
        any_src = True
        if isinstance(src, dict):
            kind = src.get("kind") or "справочник"
            if str(kind).lower() == "demo":
                kind = "Каталог платформы"
            src_data.append(
                [
                    _dash(src.get("name")),
                    _dash(src.get("url")),
                    kind,
                ]
            )
        else:
            src_data.append(
                [
                    _dash(getattr(src, "name", None)),
                    _dash(getattr(src, "url", None)),
                    "справочник",
                ]
            )
    seen_files: set[str] = set()
    for s in solutions:
        sf = s.get("source_file")
        if sf and sf not in seen_files:
            seen_files.add(sf)
            any_src = True
            src_data.append([sf, "—", _origin_label(s.get("data_origin"))])
    if not any_src:
        src_data.append(["Нет данных", "—", "—"])
    story.append(
        section(
            "Источники",
            [make_table(src_data, col_widths=[70 * mm, 55 * mm, 35 * mm])],
        )
    )

    story.append(Spacer(1, 12))
    story.append(
        p(
            "Результат является предварительной оценкой и требует верификации "
            "при обследовании объекта.",
            footer_style,
        )
    )
    story.append(
        p(
            "Источники данных указаны в разделе «Источники».",
            footer_style,
        )
    )

    def _on_page(canvas, _doc):
        canvas.saveState()
        # top accent line
        canvas.setStrokeColor(_PDF_ORANGE)
        canvas.setLineWidth(2)
        canvas.line(15 * mm, A4[1] - 10 * mm, A4[0] - 15 * mm, A4[1] - 10 * mm)
        canvas.setFont(font_name, 7.5)
        canvas.setFillColor(_PDF_MUTED)
        canvas.drawString(15 * mm, A4[1] - 8 * mm, "NETCORE")
        page_no = canvas.getPageNumber()
        canvas.drawRightString(A4[0] - 15 * mm, A4[1] - 8 * mm, f"Стр. {page_no}")
        canvas.drawCentredString(A4[0] / 2, 10 * mm, f"Страница {page_no} · предынвестиционная оценка")
        canvas.restoreState()

    doc.build(story, onFirstPage=_on_page, onLaterPages=_on_page)
    return buf.getvalue()
