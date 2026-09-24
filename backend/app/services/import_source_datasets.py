"""Импорт демо-датасетов из Датасеты_хакатон.xlsx в схемы ObjectType и демо-проекты."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.catalog_meta import ObjectType
from app.models.import_batch import ImportBatch
from app.models.misc import Normative
from app.models.project import Project, Scenario
from app.models.user import utcnow
from app.services.import_source_catalog import ImportErrorItem, ImportReport

DATASETS_FILENAME = "Датасеты_хакатон.xlsx"

DEFAULT_DATASETS_CANDIDATES = [
    Path("/data/sources") / DATASETS_FILENAME,
    Path(__file__).resolve().parents[3] / "data" / "sources" / DATASETS_FILENAME,
    Path("data/sources") / DATASETS_FILENAME,
    Path("../data/sources") / DATASETS_FILENAME,
]

SHEET_TO_OBJECT = {
    "Склад": "warehouse",
    "Аэропорт": "airport",
    "Медучреждение": "medical",
}

# Stable aliases: Russian title (lower) → existing schema key
PARAM_ALIASES: dict[str, str] = {
    "общая площадь склада": "area_m2",
    "площадь активной (роботизируемой) зоны": "active_area_m2",
    "высота потолков в зоне хранения": "ceiling_height_m",
    "количество этажей (мезонинов)": "floors_count",
    "ширина главных проездов": "main_aisle_width_m",
    "ширина рабочих проходов между стеллажами": "aisle_width_m",
    "тип напольного покрытия": "floor_type",
    "количество рабочих смен в сутки": "shifts_count",
    "продолжительность смены": "shift_hours",
    "пиковый коэффициент нагрузки": "peak_factor",
    "численность персонала склада (логистика)": "staff_count",
    "средняя годовая стоимость 1 сотрудника (с начислениями)": "staff_cost_per_person_year",
    "суммарная площадь терминала (ов)": "area_m2",
    "площадь перрона и технических зон": "apron_area_m2",
    "количество терминалов": "terminals_count",
    "количество выходов на посадку (гейтов)": "gates_count",
    "пассажиропоток (пассажиров/сутки)": "passenger_flow",
    "грузопоток (тонн/сутки)": "cargo_flow",
    "тип медицинского учреждения": "facility_type",
    "общая площадь здания(й)": "area_m2",
    "количество этажей (основной корпус)": "floors_count",
    "количество лифтов (грузовых/медицинских)": "elevators",
    "длина типичного маршрута робота": "route_length_m",
    "температура воздуха (мин)": "temperature_min_c",
    "температура воздуха (макс)": "temperature_max_c",
    "относительная влажность (макс)": "humidity_max_pct",
}

# Keys we also mirror into demo project economics helpers
MIRROR_TO_LEGACY = {
    "aisle_width_m": ("aisle_width_mm", lambda v: float(v) * 1000 if v is not None else None),
}


_TRANSLIT = {
    "а": "a",
    "б": "b",
    "в": "v",
    "г": "g",
    "д": "d",
    "е": "e",
    "ё": "e",
    "ж": "zh",
    "з": "z",
    "и": "i",
    "й": "y",
    "к": "k",
    "л": "l",
    "м": "m",
    "н": "n",
    "о": "o",
    "п": "p",
    "р": "r",
    "с": "s",
    "т": "t",
    "у": "u",
    "ф": "f",
    "х": "h",
    "ц": "ts",
    "ч": "ch",
    "ш": "sh",
    "щ": "sch",
    "ъ": "",
    "ы": "y",
    "ь": "",
    "э": "e",
    "ю": "yu",
    "я": "ya",
}


@dataclass
class DatasetParam:
    sheet: str
    row: int
    param: str
    unit: Optional[str]
    base: Any
    min_v: Any
    max_v: Any
    note: Optional[str]
    section: Optional[str] = None
    key: str = ""


@dataclass
class DatasetsImportReport:
    filename: str
    status: str = "completed"
    total_rows: int = 0
    added: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[ImportErrorItem] = field(default_factory=list)
    warnings: list[ImportErrorItem] = field(default_factory=list)
    sheets: dict[str, int] = field(default_factory=dict)
    batch_id: Optional[int] = None
    object_params_by_code: dict[str, dict[str, Any]] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "filename": self.filename,
            "status": self.status,
            "total_rows": self.total_rows,
            "added": self.added,
            "updated": self.updated,
            "skipped": self.skipped,
            "error_count": len(self.errors),
            "warning_count": len(self.warnings),
            "errors": [asdict(e) for e in self.errors[:200]],
            "warnings": [asdict(w) for w in self.warnings[:200]],
            "sheets": self.sheets,
            "batch_id": self.batch_id,
            "message": (
                f"Импорт датасетов: параметров {self.total_rows}, "
                f"обновлено схем {self.updated}, ошибок {len(self.errors)}"
            ),
        }


def resolve_datasets_path(explicit: Optional[str | Path] = None) -> Optional[Path]:
    if explicit:
        p = Path(explicit)
        return p if p.is_file() else None
    for candidate in DEFAULT_DATASETS_CANDIDATES:
        try:
            if candidate.is_file():
                return candidate.resolve()
        except OSError:
            continue
    return None


def slugify_param(name: str) -> str:
    low = (name or "").strip().lower()
    if low in PARAM_ALIASES:
        return PARAM_ALIASES[low]
    # normalize punctuation
    s = low
    for ch, lat in _TRANSLIT.items():
        s = s.replace(ch, lat)
    s = re.sub(r"[^a-z0-9]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    return (s[:64] or "param")


def _cell_str(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    if not s or s == "-":
        return None
    return s


def _cell_num(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if not isinstance(v, bool) else v
    s = str(v).strip().replace("\u00a0", " ").replace(" ", "").replace(",", ".")
    if not s or s == "-":
        return None
    try:
        return float(s)
    except ValueError:
        return str(v).strip()


def parse_dataset_sheet(ws, sheet_name: str, filename: str) -> tuple[list[DatasetParam], list[ImportErrorItem]]:
    params: list[DatasetParam] = []
    errors: list[ImportErrorItem] = []
    section: Optional[str] = None

    for row_idx in range(1, ws.max_row + 1):
        a = ws.cell(row_idx, 1).value
        b = ws.cell(row_idx, 2).value
        c = ws.cell(row_idx, 3).value
        d = ws.cell(row_idx, 4).value
        e = ws.cell(row_idx, 5).value
        f = ws.cell(row_idx, 6).value

        a_s = _cell_str(a)
        if not a_s:
            continue
        # section header
        if a_s.startswith("▌") or (b is None and c is None and "параметр" not in a_s.lower()):
            if a_s.startswith("▌") or a_s.isupper() or a_s.startswith("ДЕМО"):
                if "параметр" in a_s.lower() or a_s.startswith("ДЕМО"):
                    continue
                section = a_s
                continue
        if a_s.lower().startswith("параметр"):
            continue

        # skip if no unit and no base — likely section
        if b is None and c is None:
            section = a_s
            continue

        key = slugify_param(a_s)
        dp = DatasetParam(
            sheet=sheet_name,
            row=row_idx,
            param=a_s,
            unit=_cell_str(b),
            base=_cell_num(c),
            min_v=_cell_num(d),
            max_v=_cell_num(e),
            note=_cell_str(f),
            section=section,
            key=key,
        )
        params.append(dp)

    if not params:
        errors.append(
            ImportErrorItem(filename, sheet_name, None, None, None, "Не найдены параметры на листе")
        )
    return params, errors


def params_to_schema_properties(params: list[DatasetParam], filename: str) -> dict[str, Any]:
    props: dict[str, Any] = {}
    for p in params:
        typ = "number"
        if isinstance(p.base, str):
            typ = "string"
        prop: dict[str, Any] = {
            "title": p.param,
            "type": typ,
            "unit": p.unit,
            "default": p.base,
            "note": p.note,
            "source_file": filename,
            "source_sheet": p.sheet,
            "source_row": p.row,
        }
        if p.min_v is not None and not isinstance(p.min_v, str):
            prop["minimum"] = p.min_v
        if p.max_v is not None and not isinstance(p.max_v, str):
            prop["maximum"] = p.max_v
        if p.section:
            prop["section"] = p.section
        props[p.key] = prop
    return props


def base_params_dict(params: list[DatasetParam]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for p in params:
        if p.base is None:
            continue
        out[p.key] = p.base
        if p.key in MIRROR_TO_LEGACY:
            legacy_key, fn = MIRROR_TO_LEGACY[p.key]
            mirrored = fn(p.base)
            if mirrored is not None:
                out[legacy_key] = mirrored
    # Convenience mappings for matching/economics
    if "area_m2" in out and "area_width_m" not in out:
        # approximate square footprint for visualization
        side = (float(out["area_m2"]) ** 0.5) if out["area_m2"] else None
        if side:
            out.setdefault("area_width_m", round(side * 0.8, 1))
            out.setdefault("area_length_m", round(side * 1.25, 1))
    if "shifts_count" in out:
        shifts = out["shifts_count"]
        out.setdefault("work_mode", f"{int(shifts)}с" if shifts else "2с")
    if "staff_count" in out and "staff_cost_per_person_year" in out:
        try:
            out.setdefault(
                "current_labor_cost_year",
                float(out["staff_count"]) * float(out["staff_cost_per_person_year"]),
            )
            out.setdefault("staff_cost_year", out["current_labor_cost_year"])
        except (TypeError, ValueError):
            pass
    if "aisle_width_m" in out and "aisle_width_mm" not in out:
        try:
            out["aisle_width_mm"] = float(out["aisle_width_m"]) * 1000
        except (TypeError, ValueError):
            pass
    out.setdefault("indoor", True)
    out.setdefault("required_navigation", "SLAM")
    return out


def import_source_datasets(
    db: Session,
    path: Optional[str | Path] = None,
    content: Optional[bytes] = None,
    filename: Optional[str] = None,
    update_demo_projects: bool = True,
    commit: bool = True,
) -> DatasetsImportReport:
    resolved: Optional[Path] = None
    fname = filename or DATASETS_FILENAME
    tmp_path: Optional[Path] = None

    if content is not None:
        import tempfile

        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.write(content)
        tmp.close()
        tmp_path = Path(tmp.name)
        wb_path = tmp_path
    else:
        resolved = resolve_datasets_path(path)
        if not resolved:
            report = DatasetsImportReport(filename=fname, status="failed")
            report.errors.append(
                ImportErrorItem(fname, None, None, None, str(path) if path else "default", "Файл датасетов не найден")
            )
            return report
        wb_path = resolved
        fname = resolved.name

    report = DatasetsImportReport(filename=fname)
    try:
        wb = load_workbook(wb_path, data_only=True)
    except Exception as exc:  # noqa: BLE001
        report.status = "failed"
        report.errors.append(ImportErrorItem(fname, None, None, None, None, f"Ошибка чтения xlsx: {exc}"))
        if tmp_path:
            tmp_path.unlink(missing_ok=True)
        return report

    all_by_code: dict[str, list[DatasetParam]] = {}

    for sheet_name, obj_code in SHEET_TO_OBJECT.items():
        if sheet_name not in wb.sheetnames:
            report.warnings.append(
                ImportErrorItem(fname, sheet_name, None, None, None, "Лист отсутствует")
            )
            continue
        ws = wb[sheet_name]
        params, errs = parse_dataset_sheet(ws, sheet_name, fname)
        report.errors.extend(errs)
        report.sheets[sheet_name] = len(params)
        report.total_rows += len(params)
        all_by_code[obj_code] = params

        ot = db.query(ObjectType).filter(ObjectType.code == obj_code).first()
        if not ot:
            report.errors.append(
                ImportErrorItem(fname, sheet_name, None, None, obj_code, "ObjectType не найден")
            )
            report.skipped += len(params)
            continue

        props = params_to_schema_properties(params, fname)
        # Merge with existing schema keys to keep matching helpers
        existing = (ot.parameter_schema or {}).get("properties") or {}
        merged = {**existing, **props}
        ot.parameter_schema = {
            "type": "object",
            "properties": merged,
            "source_file": fname,
            "source_sheet": sheet_name,
        }
        report.updated += 1
        base = base_params_dict(params)
        report.object_params_by_code[obj_code] = base

        # Store normative defaults for numeric params
        for p in params:
            if p.base is None or isinstance(p.base, str):
                continue
            code = f"ds_{obj_code}_{p.key}"[:64]
            n = db.query(Normative).filter(Normative.code == code).first()
            if n:
                n.value = float(p.base)
                n.unit = p.unit
                n.source = f"{fname}:{sheet_name}:row{p.row}"
                n.name_ru = p.param
            else:
                db.add(
                    Normative(
                        code=code,
                        name_ru=p.param,
                        value=float(p.base),
                        unit=p.unit,
                        source=f"{fname}:{sheet_name}:row{p.row}",
                        editable=True,
                    )
                )
                report.added += 1

    if update_demo_projects:
        _update_demo_projects(db, report.object_params_by_code)

    report.status = "completed_with_errors" if report.errors else "completed"
    batch = ImportBatch(
        filename=fname,
        imported_at=utcnow(),
        status=report.status,
        total_rows=report.total_rows,
        added=report.added,
        updated=report.updated,
        skipped=report.skipped,
        errors=len(report.errors),
        warnings=len(report.warnings),
        report_json=report.to_dict(),
        notes="source datasets",
    )
    db.add(batch)
    if commit:
        db.commit()
        db.refresh(batch)
    else:
        db.flush()
    report.batch_id = batch.id

    if tmp_path:
        tmp_path.unlink(missing_ok=True)
    return report


DEMO_PROJECT_SPECS = [
    {
        "name": "Автоматизация склада",
        "object_code": "warehouse",
        "region": "Москва",
        "required_processes": ["transport", "picking"],
        "description": (
            "Демонстрационный проект склада. "
            "Параметры из Датасеты_хакатон.xlsx (лист «Склад»)."
        ),
    },
    {
        "name": "Автоматизация аэропортового объекта",
        "object_code": "airport",
        "region": "Санкт-Петербург",
        "required_processes": ["towing", "transport"],
        "description": (
            "Демонстрационный проект аэропорта. "
            "Параметры из Датасеты_хакатон.xlsx (лист «Аэропорт»)."
        ),
    },
    {
        "name": "Роботизация медицинского учреждения",
        "object_code": "medical",
        "region": "Казань",
        "required_processes": ["delivery", "cleaning"],
        "description": (
            "Демонстрационный проект медучреждения. "
            "Параметры из Датасеты_хакатон.xlsx (лист «Медучреждение»)."
        ),
    },
]


def _update_demo_projects(db: Session, params_by_code: dict[str, dict[str, Any]]) -> None:
    for spec in DEMO_PROJECT_SPECS:
        ot = db.query(ObjectType).filter(ObjectType.code == spec["object_code"]).first()
        if not ot:
            continue
        project = (
            db.query(Project)
            .filter(Project.name == spec["name"], Project.is_demo.is_(True))
            .first()
        )
        base = dict(params_by_code.get(spec["object_code"]) or {})
        base.setdefault("required_processes", spec["required_processes"])
        if spec["object_code"] == "airport":
            base["indoor"] = False
            base.setdefault("floor_type", "асфальт")
            base.setdefault("zone_type", "перрон")
        elif spec["object_code"] == "medical":
            base.setdefault("floor_type", "плитка")
        else:
            base.setdefault("floor_type", "бетон")

        if project:
            project.description = spec["description"]
            project.region = spec["region"]
            project.object_type_id = ot.id
            project.industry_id = ot.industry_id
            project.object_params = {**(project.object_params or {}), **base}
            project.work_mode = base.get("work_mode") or project.work_mode
            project.status = "demo"
            project.is_demo = True
        else:
            project = Project(
                owner_id=None,
                name=spec["name"],
                industry_id=ot.industry_id,
                object_type_id=ot.id,
                description=spec["description"],
                region=spec["region"],
                work_mode=base.get("work_mode") or "2с",
                status="demo",
                is_demo=True,
                object_params=base,
            )
            db.add(project)
            db.flush()
            for code, name_ru in [
                ("baseline", "Без роботизации"),
                ("purchase", "Покупка оборудования"),
                ("raas", "Роботы как услуга"),
            ]:
                exists = (
                    db.query(Scenario)
                    .filter(Scenario.project_id == project.id, Scenario.code == code)
                    .first()
                )
                if not exists:
                    db.add(
                        Scenario(
                            project_id=project.id,
                            code=code,
                            name_ru=name_ru,
                            params={},
                            results={},
                        )
                    )


def import_source_datasets_if_needed(db: Session) -> Optional[DatasetsImportReport]:
    resolved = resolve_datasets_path()
    if not resolved:
        return None
    # Re-run is idempotent (merge schemas / update demos)
    return import_source_datasets(db, path=resolved, update_demo_projects=True, commit=True)
